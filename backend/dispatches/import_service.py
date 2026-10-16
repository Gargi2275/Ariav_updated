"""Dispatch bulk import — match existing Purchase Orders only; never create a PO.

Preview is read-only. Commit re-validates and writes one Dispatch per
(Order Code, Invoice No) group inside a single transaction.
"""

from __future__ import annotations

import hashlib
import re
import secrets
from collections import OrderedDict, defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.core.cache import cache
from django.db import transaction
from django.db.models import Q
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

from accounts.access import scope_to_visible_customers
from customers.models import Customer
from purchase_orders.import_service import (
    MAX_DESCRIPTION,
    _cell_str,
    _norm_header,
    _parse_qty,
    compose_description,
)
from purchase_orders.models import PurchaseOrder, PurchaseOrderLine, PurchaseOrderStatusChange

from .models import Dispatch, DispatchLine
from .services import (
    DISPATCH_STATUS_LOCK,
    line_pending_quantity,
    qty_label,
    sync_purchase_order_dispatch_status,
)

def _exact_money(value: str) -> str | None:
    raw = _cell_str(value).replace(",", "").strip()
    if not raw or raw in {"-", "–"}:
        return None
    try:
        parsed = Decimal(raw)
    except InvalidOperation:
        return None
    if parsed != parsed.quantize(Decimal("0.01")):
        return None
    return format(parsed.quantize(Decimal("0.01")), "f")


PREVIEW_CACHE_PREFIX = "dispatch_import_preview:"
PREVIEW_TTL_SECONDS = 30 * 60

HEADER_ALIASES = {
    "companycd": "company_cd",
    "company cd": "company_cd",
    "company code": "company_cd",
    "agent": "agent",
    "agent code": "agent",
    "order code": "order_code",
    "order no": "order_code",
    "order number": "order_code",
    "invoice no": "invoice_no",
    "invoice number": "invoice_no",
    "invoice date": "invoice_date",
    "customer code": "customer_code",
    "customercd": "customer_code",
    "customer name": "customer_name",
    "customer city": "city",
    "state": "state",
    "transport name": "transporter",
    "transporter": "transporter",
    "transport": "transporter",
    "lr no": "lr_number",
    "lr number": "lr_number",
    "lr date": "lr_date",
    "code1": "code1",
    "code 1": "code1",
    "code2": "code2",
    "code 2": "code2",
    "code3": "code3",
    "code 3": "code3",
    "code4": "code4",
    "code 4": "code4",
    "code5": "code5",
    "code 5": "code5",
    "code6": "code6",
    "code 6": "code6",
    "template": "template",
    "template code": "template",
    "size": "size",
    "size code": "size",
    "blend description": "description",
    "item type": "item_type",
    "price": "price",
    "no of pcs": "pcs",
    "invoicelineno": "invoice_line_no",
    "qty": "qty",
    "quantity": "qty",
    "ord qty": "qty",
    "order qty": "qty",
    "disp qty": "qty",
    "dispatch qty": "qty",
    "meters": "qty",
    "meter": "qty",
    "product code": "product_code",
    "item code": "product_code",
    "sku": "product_code",
    "product description": "description",
    "item description": "description",
    "fabric description": "description",
    "item": "description",
}

RECORD_FIELDS = (
    "order_code",
    "invoice_no",
    "invoice_date",
    "customer_code",
    "customer_name",
    "transporter",
    "lr_number",
    "lr_date",
    "code1",
    "code2",
    "code3",
    "code4",
    "code5",
    "code6",
    "template",
    "size",
    "item_type",
    "price",
    "pcs",
    "invoice_line_no",
    "qty",
    "product_code",
    "description",
)


def group_key(order_code: str, invoice_no: str) -> str:
    return f"{order_code}||{invoice_no}"


def _parse_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
        except (OverflowError, ValueError):
            return None
    raw = _cell_str(value)
    if not raw:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(raw[:10], fmt).date()
        except ValueError:
            continue
    return None


def _line_description(record: dict) -> str:
    composed = compose_description(record)
    if composed.replace("-", "").strip():
        return composed[:MAX_DESCRIPTION]
    desc = _cell_str(record.get("description"))
    if desc:
        return desc[:MAX_DESCRIPTION]
    return _cell_str(record.get("product_code"))[:MAX_DESCRIPTION]


def _client_item_key(record: dict) -> str:
    return compose_description(record)


def _norm_match(value: str) -> str:
    return " ".join(_cell_str(value).lower().replace("-", " ").split())


def _parse_ocm_workbook(file_bytes: bytes) -> tuple[list[dict], list[str], list[dict]]:
    warnings: list[str] = []
    skipped_by_reason: OrderedDict[str, list[int]] = OrderedDict()
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError({"file": f"Could not read the spreadsheet: {exc}"}) from exc

    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            raise ValidationError({"file": "The spreadsheet is empty."}) from None

        mapping: dict[int, str] = {}
        mapped_values: set[str] = set()
        for index, cell in enumerate(header_row or ()):
            key = HEADER_ALIASES.get(_norm_header(cell))
            if not key:
                continue
            if key == "qty" and "qty" in mapped_values:
                continue
            mapping[index] = key
            mapped_values.add(key)

        missing = []
        if "order_code" not in mapped_values:
            missing.append("Order Code")
        if "invoice_no" not in mapped_values:
            missing.append("Invoice No.")
        if "qty" not in mapped_values:
            missing.append("Quantity")
        if "customer_code" not in mapped_values:
            missing.append("CustomerCd")
        if missing:
            raise ValidationError(
                {"file": f"Missing required column(s): {', '.join(missing)}."}
            )

        grouped: OrderedDict[str, dict] = OrderedDict()
        for excel_row, values in enumerate(rows, start=2):
            if not values or all(v is None or _cell_str(v) == "" for v in values):
                continue
            record = {field: None for field in RECORD_FIELDS}
            for index, field in mapping.items():
                if index < len(values):
                    record[field] = values[index]

            order_code = _cell_str(record["order_code"])
            invoice_no = _cell_str(record["invoice_no"])
            if not order_code or not invoice_no:
                reason = "Order Code or Invoice No is blank"
                skipped_by_reason.setdefault(reason, []).append(excel_row)
                warnings.append(
                    f"Row {excel_row}: skipped — Order Code or Invoice No is blank."
                )
                continue

            qty = _parse_qty(record["qty"])
            key = group_key(order_code, invoice_no)
            group = grouped.get(key)
            if group is None:
                group = {
                    "group_key": key,
                    "order_code": order_code,
                    "invoice_no": invoice_no,
                    "customer_codes": [],
                    "customer_name": _cell_str(record["customer_name"]),
                    "transporter": _cell_str(record["transporter"]),
                    "lr_number": _cell_str(record["lr_number"]),
                    "lr_date": record["lr_date"],
                    "invoice_date": record["invoice_date"],
                    "lines": [],
                    "source_rows": [],
                }
                grouped[key] = group
            group["source_rows"].append(excel_row)

            customer_code = _cell_str(record["customer_code"])
            if customer_code and customer_code not in group["customer_codes"]:
                group["customer_codes"].append(customer_code)
            if _cell_str(record["customer_name"]) and not group["customer_name"]:
                group["customer_name"] = _cell_str(record["customer_name"])
            if _cell_str(record["transporter"]) and not group["transporter"]:
                group["transporter"] = _cell_str(record["transporter"])
            if _cell_str(record["lr_number"]) and not group["lr_number"]:
                group["lr_number"] = _cell_str(record["lr_number"])
            if group["lr_date"] in (None, "") and record["lr_date"] not in (None, ""):
                group["lr_date"] = record["lr_date"]
            if group["invoice_date"] in (None, "") and record["invoice_date"] not in (None, ""):
                group["invoice_date"] = record["invoice_date"]

            if qty is None:
                reason = "Qty is missing or invalid"
                skipped_by_reason.setdefault(reason, []).append(excel_row)
                warnings.append(
                    f"Row {excel_row} (Order Code {order_code} / Invoice {invoice_no}): "
                    "skipped line — Qty is missing or invalid."
                )
                continue

            group["lines"].append(
                {
                    "description": _line_description(record),
                    "product_code": _cell_str(record["product_code"]),
                    "code1": _cell_str(record["code1"]),
                    "client_item_key": _client_item_key(record),
                    "shade": _cell_str(record["code6"]),
                    "quantity": str(_parse_qty(record["pcs"]) or qty),
                    "meters": str(qty),
                    "excel_row": excel_row,
                }
            )
    finally:
        workbook.close()

    skipped = [
        {"reason": reason, "row_numbers": row_numbers}
        for reason, row_numbers in skipped_by_reason.items()
    ]
    return list(grouped.values()), warnings, skipped


OCM_SIGNATURE_HEADERS = {"ordercode", "invoiceno"}
ARVIND_SIGNATURE_HEADERS = {"plantinvoiceline", "pono|line", "soldtoparty"}


def _header_key(value: Any) -> str:
    return " ".join(_cell_str(value).casefold().split())


def _dispatch_header_key(value: Any) -> str:
    return re.sub(r"\s+", " ", _cell_str(value).casefold()).strip()


def _read_dispatch_rows(file_bytes: bytes) -> tuple[list[list[str]], bool]:
    if file_bytes.startswith(b"\x50\x4b\x03\x04"):
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
        try:
            return [list(row) for row in workbook.active.iter_rows(values_only=True)], True
        finally:
            workbook.close()
    if file_bytes.startswith(b"\xd0\xcf\x11\xe0"):
        raise ValidationError({"file": "Legacy .xls is not supported; open it in Excel and save as .xlsx"})
    encoding = "utf-16" if file_bytes.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
    try:
        text = file_bytes.decode(encoding)
    except UnicodeDecodeError as exc:
        raise ValidationError({"file": f"Could not read the delimited dispatch file: {exc}"}) from exc
    rows = []
    for line in text.splitlines():
        if line.strip():
            rows.append([cell.strip() for cell in line.split("\t")])
    if not rows:
        raise ValidationError({"file": "The dispatch file is empty."})
    return rows, False


def detect_dispatch_format(file_bytes: bytes) -> dict:
    rows, is_xlsx = _read_dispatch_rows(file_bytes)
    headers = [_dispatch_header_key(value) for value in rows[0]]
    compact = {re.sub(r"[^a-z0-9|]+", "", value) for value in headers}
    if compact & OCM_SIGNATURE_HEADERS:
        return {"id": "ocm_dispatch", "name": "OCM Dispatch", "headers": headers, "rows": rows, "is_xlsx": is_xlsx}
    if ARVIND_SIGNATURE_HEADERS <= compact:
        return {"id": "arvind_dispatch", "name": "Arvind Dispatch", "headers": headers, "rows": rows, "is_xlsx": is_xlsx}
    po_markers = {"customercode", "customername", "customer code", "customer name"}
    if len(set(headers) & po_markers) >= 2:
        raise ValidationError({"file": "This looks like a purchase order file; use PO > Import from file"})
    known = {
        *OCM_SIGNATURE_HEADERS,
        *ARVIND_SIGNATURE_HEADERS,
        "customercode",
        "customername",
    }
    missing = sorted(known - set(compact))
    closest = "OCM Dispatch" if len(set(compact) & OCM_SIGNATURE_HEADERS) >= len(set(compact) & ARVIND_SIGNATURE_HEADERS) else "Arvind Dispatch"
    raise ValidationError({
        "file": f"No dispatch format matched. Headers found: {', '.join(_cell_str(value) for value in rows[0]) or '(none)'}. Closest profile: {closest}; missing headers: {', '.join(missing) or '(none)'}.",
    })


def _arvind_value(row: list, headers: list[str], name: str) -> str:
    aliases = {
        "invoice": {"plant invoice / line", "plant invoice/line"},
        "invoice_date": {"plant invoice date"},
        "customer_code": {"sold to party"},
        "customer_name": {"sold to party name"},
        "order": {"po no.|line", "po no. | line", "po no.| line"},
        "item": {"customer reference"},
        "pieces": {"no. of pieces", "no of pieces"},
        "meters": {"batch qty.", "batch qty"},
        "lr_number": {"lr no.", "lr no"},
        "so_price": {"so price"},
        "net_rate": {"net rate per mts", "net rate per mts."},
        "lr_date": {"lr date"},
        "transporter": {"transporter"},
        "invoice_type": {"inv. type desc.", "inv type desc."},
    }
    wanted = aliases[name]
    for index, header in enumerate(headers):
        if header in wanted:
            return _cell_str(row[index] if index < len(row) else "")
    return ""


def _arvind_date(value: str) -> tuple[date | None, str | None]:
    try:
        return datetime.strptime(value, "%d.%m.%Y").date(), None
    except ValueError:
        return None, f"Invalid date {value!r}; expected DD.MM.YYYY."


def _parse_arvind_dispatch(rows: list[list[str]]) -> tuple[list[dict], list[dict]]:
    headers = [_dispatch_header_key(value) for value in rows[0]]
    grouped: OrderedDict[str, dict] = OrderedDict()
    skipped: list[dict] = []
    for row_number, row in enumerate(rows[1:], start=2):
        invoice_raw = _arvind_value(row, headers, "invoice")
        order_raw = _arvind_value(row, headers, "order")
        invoice = invoice_raw.split("/", 1)[0].strip()
        order = order_raw.split("|", 1)[0].strip()
        invoice_type = _arvind_value(row, headers, "invoice_type")
        if "sample" in invoice_type.casefold():
            skipped.append({"reason": "Sample invoice, not a PO dispatch", "row_numbers": [row_number]})
            continue
        if not invoice or not order:
            skipped.append({"reason": "Plant Invoice / Line or PO No.|Line is blank", "row_numbers": [row_number]})
            continue
        dispatch_date, date_error = _arvind_date(_arvind_value(row, headers, "invoice_date"))
        key = group_key(order, invoice)
        group = grouped.setdefault(key, {
            "group_key": key,
            "order_code": order,
            "invoice_no": invoice,
            "customer_codes": [],
            "customer_name": _arvind_value(row, headers, "customer_name"),
            "transporter": _arvind_value(row, headers, "transporter"),
            "lr_number": _arvind_value(row, headers, "lr_number"),
            "lr_date": _arvind_value(row, headers, "lr_date"),
            "invoice_date": dispatch_date,
            "date_error": date_error,
            "lines": [],
            "source_rows": [],
            "profile": "arvind_dispatch",
        })
        group["source_rows"].append(row_number)
        customer_code = _arvind_value(row, headers, "customer_code")
        if customer_code.isdigit():
            customer_code = customer_code.zfill(10)
        if customer_code and customer_code not in group["customer_codes"]:
            group["customer_codes"].append(customer_code)
        pieces = _parse_qty(_arvind_value(row, headers, "pieces"))
        meters = _cell_str(_arvind_value(row, headers, "meters"))
        try:
            meter_value = Decimal(meters).quantize(Decimal("0.001")) if meters else None
        except InvalidOperation:
            meter_value = None
        item = _arvind_value(row, headers, "item").strip()
        if pieces is None or not item:
            skipped.append({"reason": "Customer Reference or No. of Pieces is missing or invalid", "row_numbers": [row_number]})
            continue
        group["lines"].append({
            "description": item,
            "product_code": item,
            "client_item_key": item,
            "quantity": str(pieces),
            "meters": str(meter_value) if meter_value is not None else None,
            "source_rate": _exact_money(_arvind_value(row, headers, "so_price")),
            "net_rate": _exact_money(_arvind_value(row, headers, "net_rate")),
            "excel_row": row_number,
        })
    return list(grouped.values()), skipped


def parse_workbook(file_bytes: bytes) -> tuple[list[dict], list[str], list[dict]]:
    detected = detect_dispatch_format(file_bytes)
    if detected["id"] != "ocm_dispatch":
        raise ValidationError({"file": "This file uses the Arvind dispatch format; select the Arvind brand."})
    return _parse_ocm_workbook(file_bytes)


def _po_by_number(user) -> dict[str, list[PurchaseOrder]]:
    """Index visible POs by exact brand order code and internal PO number."""
    mapping: dict[str, list[PurchaseOrder]] = defaultdict(list)
    qs = scope_to_visible_customers(PurchaseOrder.objects.all(), user).select_related("customer", "entity", "brand").prefetch_related(
        "lines__product",
        "lines__dispatch_lines",
    )
    for po in qs:
        po_number_key = po.po_number.strip().lower()
        mapping[f"po:{po_number_key}"].append(po)
        if po.brand_order_code:
            mapping[f"brand:{po.brand_order_code.strip().lower()}"].append(po)
    return mapping


def _open_customer_item_candidates(
    customer_code: str,
    item_quantities: dict[str, Decimal],
    pos: dict[str, list[PurchaseOrder]],
) -> list[PurchaseOrder]:
    candidates: list[PurchaseOrder] = []
    seen: set[int] = set()
    all_pos = {po.id: po for matches in pos.values() for po in matches}
    for po in all_pos.values():
        if po.id in seen or po.status not in {
            PurchaseOrder.Status.BRAND_ACCEPTED,
            PurchaseOrder.Status.PARTIALLY_DISPATCHED,
        }:
            continue
        if (po.customer.customer_code or "").strip().casefold() != customer_code.strip().casefold():
            continue
        matched_quantities: dict[str, Decimal] = defaultdict(lambda: Decimal("0.00"))
        for line in po.lines.all():
            key = _norm_match(line.product_description)
            if key in item_quantities:
                matched_quantities[key] += line_pending_quantity(line)
            elif line.product_id and line.product.client_item_code:
                client_key = _norm_match(line.product.client_item_code)
                if client_key in item_quantities:
                    matched_quantities[client_key] += line_pending_quantity(line)
        if all(matched_quantities.get(key, Decimal("0.00")) >= quantity for key, quantity in item_quantities.items()):
            candidates.append(po)
            seen.add(po.id)
    return candidates


def _display_item_key(value: str) -> str:
    return "-".join(_cell_str(value).replace("-", " ").split()).upper()


def _customer_item_exclusion_reason(
    customer_code: str,
    item_keys: list[str],
    pos: dict[str, list[PurchaseOrder]],
    order_code: str,
    user,
) -> str:
    all_pos = {
        po.id: po
        for matches in pos.values()
        for po in matches
        if (po.customer.customer_code or "").strip().casefold() == customer_code.strip().casefold()
    }
    visible_customers = scope_to_visible_customers(
        Customer.objects.all(), user, field="id"
    )
    customer = visible_customers.filter(
        customer_code__iexact=customer_code.strip()
    ).first()
    if customer is None:
        return f"No matching Purchase Order found for Order Code {order_code}"

    # This customer and its POs came from the scope-filtered index. Do not
    # consult the global Customer table here: doing so leaks existence and
    # diagnostics for records the operator cannot see.
    open_statuses = {
        PurchaseOrder.Status.BRAND_ACCEPTED,
        PurchaseOrder.Status.PARTIALLY_DISPATCHED,
    }
    open_pos = [po for po in all_pos.values() if po.status in open_statuses]
    ordered_keys = {
        key
        for po in open_pos
        for line in po.lines.all()
        for key in (
            _norm_match(line.product_description),
            _norm_match(line.product.client_item_code)
            if line.product_id and line.product.client_item_code
            else "",
        )
        if key
    }
    missing = ", ".join(
        dict.fromkeys(
            _display_item_key(key)
            for key in item_keys
            if _norm_match(key) not in ordered_keys
        )
    ) or "(none)"
    statuses = "; ".join(
        f"{po.po_number} is {po.status}; Brand Accepted or Partially Dispatched is required"
        for po in sorted(all_pos.values(), key=lambda item: item.po_number)
    ) or "none"
    return (
        f"No customer + item PO match for Order Code {order_code}; "
        f"customer in master: yes; "
        f"POs for this customer: {len(all_pos)}; "
        f"customer PO statuses: {statuses}; "
        f"item keys missing from customer's open POs: {missing}"
    )


def _resolve_po(order_code: str, pos: dict[str, list[PurchaseOrder]]) -> tuple[PurchaseOrder | None, str | None]:
    key = _cell_str(order_code).strip().lower()
    brand_matches = pos.get(f"brand:{key}", [])
    if len(brand_matches) > 1:
        numbers = ", ".join(po.po_number for po in brand_matches)
        return None, f"Ambiguous Order Code {order_code}; matches Purchase Orders {numbers}"
    if brand_matches:
        return brand_matches[0], None
    po_matches = pos.get(f"po:{key}", [])
    if len(po_matches) > 1:
        numbers = ", ".join(po.po_number for po in po_matches)
        return None, f"Ambiguous Order Code {order_code}; matches Purchase Orders {numbers}"
    if po_matches:
        return po_matches[0], None
    return None, None


def _order_code_exists_outside_scope(order_code: str, pos: dict[str, list[PurchaseOrder]]) -> bool:
    key = _cell_str(order_code).strip().casefold()
    visible = pos.get(f"po:{key}", []) + pos.get(f"brand:{key}", [])
    if visible:
        return False
    return PurchaseOrder.objects.filter(
        Q(po_number__iexact=key) | Q(brand_order_code__iexact=key)
    ).exists()


def _arvind_item_match(file_item: str, line: PurchaseOrderLine) -> bool:
    wanted = _cell_str(file_item).casefold().strip()
    candidates = [line.product_description]
    if line.product_id:
        candidates.extend([line.product.product_code, line.product.product_name])
    return bool(wanted) and any(wanted == _cell_str(value).casefold().strip() for value in candidates)


def _arvind_customer_item_candidates(
    group: dict,
    pos: dict[str, list[PurchaseOrder]],
) -> list[PurchaseOrder]:
    customer_code = _cell_str(group["customer_codes"][0] if group["customer_codes"] else "").casefold().strip()
    customer_name = _cell_str(group.get("customer_name")).casefold().strip()
    candidates: dict[int, PurchaseOrder] = {}
    for matches in pos.values():
        for po in matches:
            if po.brand.brand_code.casefold() != "avd":
                continue
            po_code = _cell_str(po.customer.customer_code).casefold().strip()
            po_name = _cell_str(po.customer.customer_name).casefold().strip()
            if not customer_code or po_code != customer_code:
                continue
            if customer_name and po_name != customer_name:
                continue
            if all(
                any(_arvind_item_match(source_line["client_item_key"], line) for line in po.lines.all())
                for source_line in group["lines"]
            ):
                candidates[po.id] = po
    return list(candidates.values())


def _evaluate_arvind_group(group: dict, pos: dict[str, list[PurchaseOrder]]) -> tuple[dict | None, dict | None]:
    base = {
        "group_key": group["group_key"],
        "order_code": group["order_code"],
        "invoice_no": group["invoice_no"],
        "customer_code": group["customer_codes"][0] if group["customer_codes"] else "",
        "customer_name": group["customer_name"],
        "line_count": len(group["lines"]),
        "source_rows": group["source_rows"],
    }
    if group.get("date_error"):
        return None, {**base, "reason": group["date_error"]}
    candidates = pos.get(f"brand:{group['order_code'].casefold()}", []) + pos.get(f"po:{group['order_code'].casefold()}", [])
    candidates = list({po.id: po for po in candidates}.values())
    if not candidates:
        candidates = _arvind_customer_item_candidates(group, pos)
        if not candidates:
            all_candidates = [
                po for matches in pos.values() for po in matches
                if (
                    po.brand_order_code
                    and po.brand_order_code.casefold() == group["order_code"].casefold()
                )
                or po.po_number.casefold() == group["order_code"].casefold()
            ]
            if all_candidates:
                return None, {**base, "reason": "PO belongs to another brand"}
            return None, {**base, "reason": "No matching Purchase Order"}
    arvind = [po for po in candidates if po.brand.brand_code.casefold() == "avd"]
    if not arvind:
        return None, {**base, "reason": "PO belongs to another brand"}
    if len(arvind) > 1:
        return None, {**base, "reason": f"Ambiguous Purchase Order match for {group['order_code']}"}
    po = arvind[0]
    customer_code = base["customer_code"]
    if (po.customer.customer_code or "").strip().casefold() != customer_code.casefold():
        return None, {**base, "reason": f"Customer mismatch: file Sold To Party '{customer_code}' does not match PO customer code '{po.customer.customer_code}'."}
    if Dispatch.objects.filter(purchase_order=po, challan_reference=group["invoice_no"]).exists():
        return None, {**base, "reason": f"Dispatch for {po.po_number} and invoice {group['invoice_no']} is already imported"}
    status_error = _status_reason(po)
    if status_error:
        return None, {**base, "reason": status_error}
    matched = []
    allocated: dict[int, Decimal] = defaultdict(lambda: Decimal("0.00"))
    for source_line in group["lines"]:
        matches = [line for line in po.lines.all() if _arvind_item_match(source_line["client_item_key"], line)]
        if not matches:
            return None, {**base, "reason": f"No matching PO line for item key '{source_line['client_item_key']}' on {po.po_number}."}
        if len(matches) > 1:
            return None, {**base, "reason": f"Ambiguous PO line match for '{source_line['client_item_key']}' on {po.po_number}."}
        line = matches[0]
        quantity = Decimal(source_line["quantity"])
        allocated[line.id] += quantity
        matched.append({
            "description": source_line["description"],
            "product_code": source_line["product_code"],
            "quantity": str(quantity),
            "meters": source_line["meters"],
            "source_rate": source_line.get("source_rate"),
            "net_rate": source_line.get("net_rate"),
            "purchase_order_line_id": line.id,
            "po_line_description": line.product_description,
        })
    for line_id, quantity in allocated.items():
        line = next(line for line in po.lines.all() if line.id == line_id)
        pending = line_pending_quantity(line)
        if quantity > pending:
            return None, {**base, "reason": f"Cannot dispatch {qty_label(quantity)} pieces — only {qty_label(pending)} pieces pending for this line."}
    return {
        **base,
        "po_id": po.id,
        "po_number": po.po_number,
        "po_entity_code": po.entity.short_code,
        "po_entity_name": po.entity.entity_name,
        "customer_code": po.customer.customer_code,
        "customer_name": po.customer.customer_name,
        "lr_number": group["lr_number"],
        "lr_date": group["lr_date"],
        "transporter": group["transporter"],
        "challan_reference": group["invoice_no"],
        "dispatch_date": group["invoice_date"].isoformat() if group["invoice_date"] else "",
        "opening_pending_quantity": str(sum(allocated.values())),
        "closing_pending_quantity": "0.00",
        "lines": matched,
        "format_id": "arvind_dispatch",
    }, None


def _file_keys(line: dict) -> set[str]:
    keys: set[str] = set()
    client_key = _norm_match(line.get("client_item_key") or "")
    if client_key:
        keys.add(client_key)
    code1 = _norm_match(line.get("code1") or "")
    if code1:
        keys.add(code1)
    desc = _norm_match(line.get("description") or "")
    if desc:
        keys.add(desc)
    code = _norm_match(line.get("product_code") or "")
    if code:
        keys.add(code)
    return keys


def _po_line_keys(line: PurchaseOrderLine) -> set[str]:
    keys: set[str] = set()
    if line.product_id and line.product.client_item_code:
        client_key = _norm_match(line.product.client_item_code)
        if client_key:
            keys.add(client_key)
    desc = _norm_match(line.product_description)
    if desc:
        keys.add(desc)
    if line.product_id:
        code = _norm_match(line.product.product_code)
        name = _norm_match(line.product.product_name)
        if code:
            keys.add(code)
        if name:
            keys.add(name)
    return keys


def match_lines_to_po(group_lines: list[dict], po: PurchaseOrder) -> tuple[list[dict], str | None]:
    """Return (allocated lines, error reason). Error means exclude the group."""
    po_lines = list(po.lines.all())
    if not po_lines:
        return [], f"Purchase Order {po.po_number} has no line items to dispatch against."

    allocated: dict[int, Decimal] = defaultdict(lambda: Decimal("0.00"))
    matched_rows: list[dict] = []

    for file_line in group_lines:
        qty = Decimal(str(file_line["quantity"]))
        item_key = _norm_match(file_line.get("client_item_key") or "")
        if item_key:
            candidates = [
                line for line in po_lines
                if item_key == _norm_match(line.product_description)
                or (
                    line.product_id
                    and line.product.client_item_code
                    and item_key == _norm_match(line.product.client_item_code)
                )
            ]
            if not candidates:
                file_keys = _file_keys(file_line)
                candidates = [
                    line for line in po_lines if file_keys and (file_keys & _po_line_keys(line))
                ]
        else:
            file_keys = _file_keys(file_line)
            candidates = [
                line for line in po_lines if file_keys and (file_keys & _po_line_keys(line))
            ]
        unique_ids = {line.id for line in candidates}
        if not unique_ids:
            label = file_line.get("description") or file_line.get("product_code") or "(blank)"
            return [], f"No matching PO line for item key '{item_key or label}' on {po.po_number}."
        if len(unique_ids) > 1:
            label = file_line.get("description") or file_line.get("product_code") or "(blank)"
            return [], (
                f"Ambiguous PO line match for '{label}' on {po.po_number} "
                f"— {len(unique_ids)} lines could match. Resolve on the PO before importing."
            )
        po_line = candidates[0]
        allocated[po_line.id] += qty
        matched_rows.append(
            {
                "description": file_line.get("description") or "",
                "product_code": file_line.get("product_code") or "",
                "code1": file_line.get("code1") or "",
                "client_item_key": file_line.get("client_item_key") or "",
                "shade": file_line.get("shade") or "",
                "quantity": str(qty),
                "purchase_order_line_id": po_line.id,
                "po_line_description": po_line.product_description
                or (po_line.product.product_name if po_line.product_id else ""),
            }
        )

    by_id = {line.id: line for line in po_lines}
    for po_line_id, extra in allocated.items():
        po_line = by_id[po_line_id]
        pending = line_pending_quantity(po_line)
        if extra > pending:
            label = (
                po_line.product_description
                or (po_line.product.product_code if po_line.product_id else f"line {po_line_id}")
            )
            return [], (
                f"Cannot dispatch {qty_label(extra)} units — only "
                f"{qty_label(pending)} units pending for this line ({label})."
            )

    collapsed: dict[int, dict] = {}
    for row in matched_rows:
        line_id = row["purchase_order_line_id"]
        if line_id not in collapsed:
            collapsed[line_id] = dict(row)
        else:
            collapsed[line_id]["quantity"] = str(
                Decimal(collapsed[line_id]["quantity"]) + Decimal(row["quantity"])
            )
    return list(collapsed.values()), None


def _status_reason(po: PurchaseOrder) -> str | None:
    if po.status in {
        PurchaseOrder.Status.BRAND_ACCEPTED,
        PurchaseOrder.Status.PARTIALLY_DISPATCHED,
    }:
        return None
    return f"PO status must be Brand Accepted or Partially Dispatched (current status: {po.status})"


def _customer_mismatch_reason(po: PurchaseOrder, file_code: str) -> str:
    actual = (po.customer.customer_code or "").strip()
    return (
        f"Customer mismatch: file has Customer Code {file_code or '(blank)'} "
        f"but Purchase Order {po.po_number} belongs to {actual} "
        f"({po.customer.customer_name})."
    )


def evaluate_group(group: dict, pos: dict[str, list[PurchaseOrder]], user) -> tuple[dict | None, dict | None]:
    """Return (importable_row, excluded_row) — exactly one is set."""
    order_code = group["order_code"]
    invoice_no = group["invoice_no"]
    codes = [c for c in group["customer_codes"] if c]
    customer_code = codes[0] if codes else ""
    customer_name = group["customer_name"]
    line_count = len(group["lines"])
    base_excluded = {
        "group_key": group["group_key"],
        "order_code": order_code,
        "invoice_no": invoice_no,
        "customer_code": customer_code,
        "customer_name": customer_name,
        "line_count": line_count,
        "source_rows": group["source_rows"],
    }

    if line_count == 0:
        return None, {**base_excluded, "reason": "No valid line items"}

    if len(codes) > 1:
        return None, {
            **base_excluded,
            "customer_code": ", ".join(codes),
            "reason": "Invoice/order group has mixed customer codes",
        }

    po, resolve_err = _resolve_po(order_code, pos)
    matched_by_customer_item = False
    if resolve_err:
        return None, {**base_excluded, "reason": resolve_err}
    if po is None:
        if _order_code_exists_outside_scope(order_code, pos):
            return None, {
                **base_excluded,
                "reason": f"No matching Purchase Order found for Order Code {order_code}",
            }
        item_quantities: dict[str, Decimal] = defaultdict(lambda: Decimal("0.00"))
        for line in group["lines"]:
            item_quantities[_norm_match(line["client_item_key"])] += Decimal(line["quantity"])
        candidates = _open_customer_item_candidates(customer_code, item_quantities, pos)
        if len(candidates) > 1:
            numbers = ", ".join(sorted(candidate.po_number for candidate in candidates))
            return None, {**base_excluded, "reason": f"Ambiguous customer + item match; candidate POs: {numbers}"}
        if not candidates:
            return None, {
                **base_excluded,
                "reason": _customer_item_exclusion_reason(
                    customer_code,
                    list(item_quantities),
                    pos,
                    order_code,
                    user=user,
                ),
            }
        po = candidates[0]
        matched_by_customer_item = True
    base_excluded.update(
        {
            "po_id": po.id,
            "po_number": po.po_number,
            "po_entity_code": po.entity.short_code,
            "po_entity_name": po.entity.entity_name,
        }
    )

    po_code = (po.customer.customer_code or "").strip()
    if not customer_code or customer_code.strip().casefold() != po_code.casefold():
        return None, {
            **base_excluded,
            "reason": (
                f"Customer mismatch: file CustomerCd '{customer_code}' "
                f"does not match PO customer code '{po_code}' for {po.po_number}."
            ),
        }

    if Dispatch.objects.filter(purchase_order=po, challan_reference=invoice_no).exists():
        return None, {**base_excluded, "reason": f"Dispatch for {po.po_number} and invoice {invoice_no} is already imported"}

    status_err = _status_reason(po)
    if status_err:
        return None, {**base_excluded, "reason": status_err}

    lr_date = _parse_date(group["lr_date"]) or _parse_date(group["invoice_date"])
    if lr_date is None:
        return None, {**base_excluded, "reason": "LR Date is missing or invalid"}

    lr_number = group["lr_number"]
    if not lr_number:
        return None, {**base_excluded, "reason": "LR No is missing"}

    matched_lines, match_err = match_lines_to_po(group["lines"], po)
    if match_err:
        return None, {**base_excluded, "reason": match_err}
    opening_quantity = sum(
        (line_pending_quantity(po_line) for po_line in po.lines.all()),
        Decimal("0.00"),
    )
    imported_quantity = sum(
        (Decimal(row["quantity"]) for row in matched_lines),
        Decimal("0.00"),
    )

    return (
        {
            "group_key": group["group_key"],
            "order_code": order_code,
            "invoice_no": invoice_no,
            "po_id": po.id,
            "po_number": po.po_number,
            "po_entity_code": po.entity.short_code,
            "po_entity_name": po.entity.entity_name,
            "customer_code": po.customer.customer_code,
            "customer_name": po.customer.customer_name,
            "file_customer_code": customer_code,
            "matched_by_customer_item": matched_by_customer_item,
            "lr_number": lr_number,
            "lr_date": lr_date.isoformat(),
            "transporter": group["transporter"],
            "challan_reference": invoice_no,
            "line_count": len(group["lines"]),
            "opening_pending_quantity": str(opening_quantity),
            "closing_pending_quantity": str(opening_quantity - imported_quantity),
            "lines": matched_lines,
        },
        None,
    )


def build_preview(file_bytes: bytes, *, user) -> dict:
    detected = detect_dispatch_format(file_bytes)
    if detected["id"] == "arvind_dispatch":
        groups, skipped = _parse_arvind_dispatch(detected["rows"])
        pos = _po_by_number(user)
        importable: list[dict] = []
        excluded: list[dict] = []
        for skipped_row in skipped:
            for row_number in skipped_row["row_numbers"]:
                excluded.append({
                    "group_key": f"row-{row_number}",
                    "order_code": "",
                    "invoice_no": "",
                    "customer_code": "",
                    "customer_name": "",
                    "reason": skipped_row["reason"],
                    "line_count": 1,
                    "source_rows": [row_number],
                })
        for group in groups:
            row, skip = _evaluate_arvind_group(group, pos)
            if skip:
                excluded.append(skip)
            else:
                importable.append(row)
        return {
            "file_sha256": hashlib.sha256(file_bytes).hexdigest(),
            "format_id": detected["id"],
            "format_name": detected["name"],
            "importable": importable,
            "excluded": excluded,
            "warnings": [],
            "skipped": [],
            "counts": {
                "rows": len(detected["rows"]) - 1,
                "importable": sum(row["line_count"] for row in importable),
                "sample": sum(len(item["source_rows"]) for item in excluded if item["reason"] == "Sample invoice, not a PO dispatch"),
                "excluded": sum(len(item["source_rows"]) for item in excluded if item["reason"] != "Sample invoice, not a PO dispatch"),
            },
        }
    groups, parse_warnings, skipped = _parse_ocm_workbook(file_bytes)
    pos = _po_by_number(user)
    importable: list[dict] = []
    excluded: list[dict] = []
    warnings = list(parse_warnings)

    for group in groups:
        row, skip = evaluate_group(group, pos, user)
        if skip:
            excluded.append(skip)
        else:
            importable.append(row)

    return {
        "file_sha256": hashlib.sha256(file_bytes).hexdigest(),
        "format_id": "ocm_dispatch",
        "format_name": "OCM Dispatch",
        "importable": importable,
        "excluded": excluded,
        "warnings": warnings,
        "skipped": skipped,
        "counts": {
            "rows": sum(len(group["source_rows"]) for group in groups) + sum(len(item["row_numbers"]) for item in skipped),
            "importable": sum(row["line_count"] for row in importable),
            "sample": 0,
            "excluded": sum(row["line_count"] for row in excluded),
        },
    }


def store_preview(user_id: int, preview: dict) -> str:
    token = secrets.token_urlsafe(32)
    cache.set(
        f"{PREVIEW_CACHE_PREFIX}{token}",
        {"user_id": user_id, **preview},
        timeout=PREVIEW_TTL_SECONDS,
    )
    return token


def load_preview(token: str) -> dict | None:
    if not token:
        return None
    cached = cache.get(f"{PREVIEW_CACHE_PREFIX}{token}")
    return cached if isinstance(cached, dict) else None


def drop_preview(token: str) -> None:
    if token:
        cache.delete(f"{PREVIEW_CACHE_PREFIX}{token}")


def public_preview_payload(preview: dict, token: str) -> dict:
    importable = []
    for row in preview["importable"]:
        item = {
                "group_key": row["group_key"],
                "order_code": row["order_code"],
                "invoice_no": row["invoice_no"],
                "po_id": row["po_id"],
                "po_number": row["po_number"],
                "po_entity_code": row["po_entity_code"],
                "po_entity_name": row["po_entity_name"],
                "match_message": (
                    f"matched by customer + item; brand order code {row['order_code']} will be recorded"
                    if row.get("matched_by_customer_item")
                    else ""
                ),
                "customer_code": row["customer_code"],
                "customer_name": row["customer_name"],
                "lr_number": row["lr_number"],
                "lr_date": row["lr_date"],
                "transporter": row["transporter"],
                "challan_reference": row["challan_reference"],
                "line_count": row["line_count"],
                "opening_pending_quantity": row["opening_pending_quantity"],
                "closing_pending_quantity": row["closing_pending_quantity"],
                "lines": row["lines"],
            }
        if row.get("format_id") == "arvind_dispatch":
            item["dispatch_date"] = row["dispatch_date"]
            for line in item["lines"]:
                line["meters"] = line.get("meters")
        importable.append(item)
    payload = {
        "preview_token": token,
        "importable": importable,
        "excluded": preview["excluded"],
        "warnings": preview["warnings"],
        "skipped": preview["skipped"],
    }
    for key in ("format_id", "format_name", "counts"):
        if key in preview:
            payload[key] = preview[key]
    return payload


def _revalidate_importable(row: dict, pos: dict[str, list[PurchaseOrder]]) -> str | None:
    if row.get("matched_by_customer_item"):
        po = next(
            (
                candidate
                for matches in pos.values()
                for candidate in matches
                if candidate.id == row["po_id"]
            ),
            None,
        )
        resolve_err = None
    else:
        po, resolve_err = _resolve_po(row["order_code"], pos)
    if resolve_err:
        return resolve_err
    if po is None:
        return f"No matching Purchase Order found for Order Code {row['order_code']}"
    if po.id != row["po_id"]:
        return f"Purchase Order match for {row['order_code']} changed since preview. Run preview again."
    file_code = (row.get("file_customer_code") or "").strip()
    po_code = (po.customer.customer_code or "").strip()
    if not file_code or file_code.casefold() != po_code.casefold():
        return f"Customer mismatch: file CustomerCd '{file_code}' does not match PO customer code '{po_code}' for {po.po_number}."
    if Dispatch.objects.filter(purchase_order=po, challan_reference=row["invoice_no"]).exists():
        return f"Dispatch for {po.po_number} and invoice {row['invoice_no']} is already imported"
    status_err = _status_reason(po)
    if status_err:
        return status_err
    file_lines = [
        {
            "description": line["description"],
            "product_code": line.get("product_code") or "",
            "code1": line.get("code1") or "",
            "client_item_key": line.get("client_item_key") or "",
            "shade": line.get("shade") or "",
            "quantity": line["quantity"],
        }
        for line in row["lines"]
    ]
    _matched, match_err = match_lines_to_po(file_lines, po)
    return match_err


def _commit_arvind_import(*, user, token: str, cached: dict, to_create: list[dict], skipped_excluded: list[dict]) -> dict:
    created: list[dict] = []
    with transaction.atomic():
        for row in to_create:
            po = PurchaseOrder.objects.select_for_update().select_related("brand", "customer").get(id=row["po_id"])
            if po.brand.brand_code.casefold() != "avd":
                raise ValidationError({"detail": "PO belongs to another brand"})
            if Dispatch.objects.filter(purchase_order=po, challan_reference=row["invoice_no"]).exists():
                raise ValidationError({"detail": f"Dispatch for {po.po_number} and invoice {row['invoice_no']} is already imported"})
            status_err = _status_reason(po)
            if status_err:
                raise ValidationError({"detail": status_err})
            dispatch = Dispatch.objects.create(
                purchase_order=po,
                dispatch_date=date.fromisoformat(row["dispatch_date"]),
                lr_number=row.get("lr_number", ""),
                transporter=row.get("transporter", ""),
                challan_reference=row["challan_reference"],
                created_by=user if getattr(user, "is_authenticated", False) else None,
            )
            for line in row["lines"]:
                DispatchLine.objects.create(
                    dispatch=dispatch,
                    purchase_order_line_id=line["purchase_order_line_id"],
                    dispatched_quantity=Decimal(str(line["quantity"])),
                    dispatched_meters=Decimal(str(line["meters"])) if line.get("meters") not in (None, "") else None,
                    source_rate=Decimal(str(line["source_rate"])) if line.get("source_rate") not in (None, "") else None,
                    net_rate=Decimal(str(line["net_rate"])) if line.get("net_rate") not in (None, "") else None,
                )
            sync_purchase_order_dispatch_status(po, user=user, source=PurchaseOrderStatusChange.Source.IMPORT)
            po.refresh_from_db()
            from notifications.services import notify_dispatch_created

            notify_dispatch_created(dispatch)
            created.append(
                {
                    "id": dispatch.id,
                    "po_id": po.id,
                    "po_number": po.po_number,
                    "po_status": po.status,
                    "lr_number": dispatch.lr_number,
                    "invoice_no": row["invoice_no"],
                    "customer_name": po.customer.customer_name,
                    "line_count": len(row["lines"]),
                }
            )
    drop_preview(token)
    remaining_excluded = list(cached["excluded"])
    return {
        "created": created,
        "created_count": len(created),
        "excluded": remaining_excluded,
        "excluded_count": len(remaining_excluded),
        "skipped_excluded_group_keys": [row["group_key"] for row in skipped_excluded],
    }


def commit_import(*, user, token: str, group_keys: list[str]) -> dict:
    cached = load_preview(token)
    if not cached:
        raise ValidationError(
            {"detail": "Preview has expired or was not found. Run preview again before committing."}
        )
    if cached.get("user_id") != getattr(user, "id", None):
        raise ValidationError(
            {"detail": "Preview has expired or was not found. Run preview again before committing."}
        )

    approved = [_cell_str(value) for value in group_keys if _cell_str(value)]
    if not approved:
        raise ValidationError({"group_keys": "Select at least one dispatch group to import."})

    importable_by_key = {row["group_key"]: row for row in cached["importable"]}
    excluded_by_key = {row["group_key"]: row for row in cached["excluded"]}
    unknown = [key for key in approved if key not in importable_by_key and key not in excluded_by_key]
    if unknown:
        raise ValidationError(
            {
                "group_keys": (
                    "Group "
                    + ", ".join(unknown)
                    + " was not part of the preview. Run preview again."
                )
            }
        )

    to_create = [importable_by_key[key] for key in approved if key in importable_by_key]
    skipped_excluded = [excluded_by_key[key] for key in approved if key in excluded_by_key]

    if cached.get("format_id") == "arvind_dispatch":
        return _commit_arvind_import(
            user=user,
            token=token,
            cached=cached,
            to_create=to_create,
            skipped_excluded=skipped_excluded,
        )

    pos = _po_by_number(user)
    stale: list[str] = []
    for row in to_create:
        err = _revalidate_importable(row, pos)
        if err:
            stale.append(err)
    if stale:
        raise ValidationError(
            {
                "detail": (
                    "Import was rejected because data changed since preview. "
                    + " ".join(stale)
                )
            }
        )

    created: list[dict] = []
    with transaction.atomic():
        for row in to_create:
            if row.get("matched_by_customer_item"):
                po = next(
                    (
                        candidate
                        for matches in pos.values()
                        for candidate in matches
                        if candidate.id == row["po_id"]
                    ),
                    None,
                )
                resolve_err = None
            else:
                po, resolve_err = _resolve_po(row["order_code"], pos)
            if resolve_err or po is None:
                raise ValidationError({"detail": resolve_err or f"No matching Purchase Order found for Order Code {row['order_code']}"})
            if row.get("matched_by_customer_item") and not po.brand_order_code:
                po.brand_order_code = row["order_code"].strip()
                po.save(update_fields=["brand_order_code", "updated_at"])
            file_lines = [
                {
                    "description": line["description"],
                    "product_code": line.get("product_code") or "",
                    "code1": line.get("code1") or "",
                    "client_item_key": line.get("client_item_key") or "",
                    "shade": line.get("shade") or "",
                    "quantity": line["quantity"],
                }
                for line in row["lines"]
            ]
            matched, match_err = match_lines_to_po(file_lines, po)
            if match_err:
                raise ValidationError({"detail": match_err})

            dispatch = Dispatch.objects.create(
                purchase_order=po,
                dispatch_date=date.fromisoformat(row["lr_date"]),
                lr_number=row["lr_number"],
                transporter=row["transporter"],
                challan_reference=row["challan_reference"],
                created_by=user if getattr(user, "is_authenticated", False) else None,
            )
            for line in matched:
                DispatchLine.objects.create(
                    dispatch=dispatch,
                    purchase_order_line_id=line["purchase_order_line_id"],
                    dispatched_quantity=Decimal(str(line["quantity"])),
                )
            sync_purchase_order_dispatch_status(po, user=user, source=PurchaseOrderStatusChange.Source.IMPORT)
            po.refresh_from_db()
            from notifications.services import notify_dispatch_created

            notify_dispatch_created(dispatch)
            created.append(
                {
                    "id": dispatch.id,
                    "po_id": po.id,
                    "po_number": po.po_number,
                    "po_status": po.status,
                    "lr_number": dispatch.lr_number,
                    "invoice_no": row["invoice_no"],
                    "customer_name": po.customer.customer_name,
                    "line_count": len(matched),
                }
            )
            pos = _po_by_number(user)

    drop_preview(token)
    remaining_excluded = list(cached["excluded"])
    return {
        "created": created,
        "created_count": len(created),
        "excluded": remaining_excluded,
        "excluded_count": len(remaining_excluded),
        "skipped_excluded_group_keys": [row["group_key"] for row in skipped_excluded],
    }
