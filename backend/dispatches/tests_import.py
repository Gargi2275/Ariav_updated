from decimal import Decimal
from io import BytesIO
from collections import defaultdict

from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook
from rest_framework.test import APIClient

from django.test import TestCase

from accounts.models import AuthUser
from brands.models import Brand
from categories.models import Category
from customers.models import Customer, CustomerEntity
from dispatches.models import Dispatch
from dispatches.models import DispatchLine
from entities.models import Entity
from products.models import Product
from purchase_orders.models import PurchaseOrder
from purchase_orders.models import PurchaseOrderLine
from dispatches.import_service import parse_workbook
from purchase_orders.import_service import parse_workbook as parse_po_workbook


HEADERS = [
    "CompanyCd",
    "Agent",
    "Order Code",
    "Invoice No",
    "Invoice Date",
    "Customer Code",
    "Customer Name",
    "Customer City",
    "State",
    "Code1",
    "Code5",
    "Code6",
    "Template",
    "Size",
    "Qty",
    "Transport Name",
    "LR No",
    "LR Date",
]


def _xlsx(rows: list[list]) -> SimpleUploadedFile:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(
        "aditi_and_shriva_dispatch_ocm.xlsx",
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _row(
    order_code,
    invoice_no,
    customer_code="DSP-C1",
    customer_name="Dispatch Trading",
    qty=4,
    code1="A",
    code5="B",
    code6="C",
    template="T",
    size="S",
    lr_no="LR-1",
    lr_date="2026-09-17",
    transport="VRL Logistics",
    invoice_date="2026-09-17",
):
    return [
        "ADT",
        "AG1",
        order_code,
        invoice_no,
        invoice_date,
        customer_code,
        customer_name,
        "Surat",
        "Gujarat",
        code1,
        code5,
        code6,
        template,
        size,
        qty,
        transport,
        lr_no,
        lr_date,
    ]


class DispatchImportApiTests(TestCase):
    def setUp(self):
        cache.clear()
        self.admin = AuthUser.objects.create_user(
            username="dsp-imp-admin",
            password="pass12345",
            role=AuthUser.Role.ADMIN,
        )
        self.operator = AuthUser.objects.create_user(
            username="dsp-imp-op",
            password="pass12345",
            role=AuthUser.Role.OPERATOR,
        )
        self.client = APIClient()
        self.client.force_authenticate(self.admin)
        self.entity = Entity.objects.create(
            short_code="ADT",
            entity_name="Aditi",
            entity_type=Entity.EntityType.HEAD_OFFICE,
            address_line_1="Aditi Road",
            city="Surat",
            state="Gujarat",
        )
        self.customer = Customer.objects.create(
            customer_code="DSP-C1",
            customer_name="Dispatch Trading",
        )
        self.other = Customer.objects.create(
            customer_code="DSP-C2",
            customer_name="Other House",
        )
        CustomerEntity.objects.create(
            customer=self.customer,
            entity=self.entity,
            primary_entity=True,
        )
        self.brand = Brand.objects.create(
            brand_code="OCM",
            brand_name="OCM",
            order_method=Brand.OrderMethod.MANUAL_POR,
        )
        root = Category.objects.create(category_code="YRN", category_name="Yarn")
        sub = Category.objects.create(
            category_code="CTN",
            category_name="Cotton",
            parent_category=root,
        )
        self.product = Product.objects.create(
            product_code="DSP-01",
            product_name="60s Cotton",
            brand=self.brand,
            category=sub,
            rate=Decimal("10.00"),
            unit=Product.Unit.METER,
        )

    def _create_po(
        self,
        po_number="ATN24R00001",
        status="Draft",
        qty="10.00",
        description="A-B-C-T-S",
        brand_order_code=None,
        brand_id=None,
    ):
        res = self.client.post(
            "/api/purchase-orders/",
            {
                "po_number": po_number,
                "brand_order_code": brand_order_code,
                "entity_id": self.entity.id,
                "customer_id": self.customer.id,
                "brand_id": brand_id or self.brand.id,
                "po_date": "2026-09-16",
                "lines": [
                    {
                        "product_id": None,
                        "product_description": description,
                        "quantity": qty,
                        "rate": "10.00",
                    }
                ],
            },
            format="json",
            HTTP_HOST="localhost",
        )
        self.assertEqual(res.status_code, 201, res.content)
        po_id = res.json()["id"]
        if status != "Draft":
            self.client.post(f"/api/purchase-orders/{po_id}/submit/", HTTP_HOST="localhost")
            path = [
                "Submitted",
                "Sent to Brand",
                "Brand Accepted",
                "Partially Dispatched",
                "Fully Dispatched",
            ]
            current = "Submitted"
            for nxt in path[1:]:
                if current == status:
                    break
                step = self.client.post(
                    f"/api/purchase-orders/{po_id}/status/",
                    {"status": nxt},
                    format="json",
                    HTTP_HOST="localhost",
                )
                self.assertEqual(step.status_code, 200, step.content)
                current = nxt
                if current == status:
                    break
        return res.json()["id"]

    def test_client_fixture_has_expected_shape(self):
        with open(
            "dispatches/tests/fixtures/aditi_and_shriva_dispatch_ocm.xlsx",
            "rb",
        ) as source:
            groups, warnings, skipped = parse_workbook(source.read())
        self.assertFalse(warnings)
        self.assertFalse(skipped)
        self.assertEqual(sum(len(group["lines"]) for group in groups), 38)
        self.assertEqual(len(groups), 12)
        self.assertEqual(
            groups[0]["lines"][0]["client_item_key"],
            "TV-0292200-163-P01-FS",
        )

    def test_preview_matches_brand_order_code_before_po_number(self):
        self._create_po("INTERNAL-PO", status="Brand Accepted", brand_order_code="CLIENT-PO")
        res = self._preview([_row("CLIENT-PO", "INV-BRAND")])
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(res.json()["importable"][0]["po_number"], "INTERNAL-PO")

    def test_preview_rejects_ambiguous_brand_order_code(self):
        other_brand = Brand.objects.create(
            brand_code="OCM2",
            brand_name="OCM 2",
            order_method=Brand.OrderMethod.MANUAL_POR,
        )
        self._create_po("INTERNAL-A", status="Brand Accepted", brand_order_code="CLIENT-DUP")
        self._create_po("INTERNAL-B", status="Brand Accepted", brand_order_code="CLIENT-DUP", brand_id=other_brand.id)
        res = self._preview([_row("CLIENT-DUP", "INV-DUP")])
        self.assertEqual(res.status_code, 200, res.content)
        self.assertIn("Ambiguous Order Code", res.json()["excluded"][0]["reason"])

    def _preview(self, rows, auth=None):
        if auth:
            self.client.force_authenticate(auth)
        upload = _xlsx(rows)
        res = self.client.post(
            "/api/dispatches/import/preview/",
            {"file": upload},
            format="multipart",
            HTTP_HOST="localhost",
        )
        return res

    def _fixture_bytes(self):
        with open(
            "dispatches/tests/fixtures/aditi_and_shriva_dispatch_ocm.xlsx",
            "rb",
        ) as source:
            return source.read()

    def _fixture_groups(self):
        groups, warnings, skipped = parse_workbook(self._fixture_bytes())
        self.assertFalse(warnings)
        self.assertFalse(skipped)
        return groups

    def _create_fixture_pos(self):
        groups = self._fixture_groups()
        customers = {}
        products = {}
        quantities = defaultdict(Decimal)
        groups_by_order = defaultdict(list)
        for group in groups:
            groups_by_order[group["order_code"]].append(group)
            customer_code = group["customer_codes"][0]
            if customer_code not in customers:
                customer = Customer.objects.create(
                    customer_code=customer_code,
                    customer_name=f"Fixture Customer {customer_code}",
                )
                CustomerEntity.objects.create(
                    customer=customer,
                    entity=self.entity,
                    primary_entity=True,
                )
                customers[customer_code] = customer
            for line in group["lines"]:
                key = line["client_item_key"]
                quantities[(group["order_code"], key)] += Decimal(line["quantity"])
                if key not in products:
                    products[key] = Product.objects.create(
                        product_code=f"CLIENT-{len(products) + 1}",
                        product_name=line["description"],
                        client_item_code=key,
                        brand=self.brand,
                        category=self.product.category,
                        rate=Decimal("10.00"),
                        unit=Product.Unit.METER,
                    )

        po_ids = {}
        for index, (order_code, order_groups) in enumerate(groups_by_order.items()):
            group = order_groups[0]
            po = PurchaseOrder.objects.create(
                po_number=f"FIXTURE-INTERNAL-{index + 1}",
                brand_order_code=order_code,
                entity=self.entity,
                customer=customers[group["customer_codes"][0]],
                brand=self.brand,
                po_date="2026-09-16",
                status=PurchaseOrder.Status.BRAND_ACCEPTED,
            )
            keys = {
                line["client_item_key"]
                for order_group in order_groups
                for line in order_group["lines"]
            }
            for key in keys:
                PurchaseOrderLine.objects.create(
                    purchase_order=po,
                    product=products[key],
                    quantity=quantities[(order_code, key)] + Decimal("100.00"),
                    rate=Decimal("10.00"),
                    line_total=Decimal("10.00"),
                )
            po_ids[order_code] = po.id
        return po_ids, customers

    def _fixture_upload(self):
        return SimpleUploadedFile(
            "aditi_and_shriva_dispatch_ocm.xlsx",
            self._fixture_bytes(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def _po_fixture_bytes(self):
        with open("dispatches/tests/fixtures/OCM_PO.xlsx", "rb") as source:
            return source.read()

    def _import_po_fixture(self):
        po_groups, warnings = parse_po_workbook(self._po_fixture_bytes())
        self.assertFalse(warnings)
        for code in {
            code
            for group in po_groups
            for code in group["customer_codes"]
        }:
            customer = Customer.objects.create(
                customer_code=code,
                customer_name=f"PO Fixture Customer {code}",
            )
            CustomerEntity.objects.create(
                customer=customer,
                entity=self.entity,
                primary_entity=True,
            )
        preview = self.client.post(
            "/api/purchase-orders/import/preview/",
            {
                "file": SimpleUploadedFile(
                    "OCM_PO.xlsx",
                    self._po_fixture_bytes(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
                "entity_id": self.entity.id,
                "brand_id": self.brand.id,
            },
            format="multipart",
            HTTP_HOST="localhost",
        )
        self.assertEqual(preview.status_code, 200, preview.content)
        body = preview.json()
        self.assertEqual(len(body["importable"]), 10)
        commit = self.client.post(
            "/api/purchase-orders/import/commit/",
            {
                "preview_token": body["preview_token"],
                "entity_id": self.entity.id,
                "brand_id": self.brand.id,
                "order_nos": [row["order_no"] for row in body["importable"]],
            },
            format="json",
            HTTP_HOST="localhost",
        )
        self.assertEqual(commit.status_code, 201, commit.content)
        PurchaseOrder.objects.all().update(status=PurchaseOrder.Status.BRAND_ACCEPTED)

    def _dispatch_fixture_preview(self):
        return self.client.post(
            "/api/dispatches/import/preview/",
            {"file": self._fixture_upload()},
            format="multipart",
            HTTP_HOST="localhost",
        )

    def test_real_po_import_and_dispatch_customer_item_fallback(self):
        self._import_po_fixture()
        first_po = PurchaseOrder.objects.order_by("id").first()
        self.assertIsNotNone(first_po)
        first_group = self._fixture_groups()[0]
        first_po.brand_order_code = ""
        first_po.save()
        response = self._dispatch_fixture_preview()
        self.assertEqual(response.status_code, 200, response.content)
        body = response.json()
        self.assertEqual(body["importable"], [])
        self.assertEqual(len(body["excluded"]), 12)

    def test_customer_item_match_uses_pieces_and_records_brand_order_code(self):
        po_id = self._create_po("CUSTOMER-ITEM-PO", qty="5.00", description="A-B-C-T-S")
        po = PurchaseOrder.objects.get(pk=po_id)
        po.brand_order_code = None
        po.status = PurchaseOrder.Status.BRAND_ACCEPTED
        po.save()
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Order Code", "Invoice No.", "CustomerCd", "Quantity", "No OF PCS",
            "Code1", "Code5", "Code6", "Template Code", "Size Code", "LR No.", "LR Date",
        ])
        sheet.append([
            "ATS-CUSTOMER-ITEM", "INV-PIECES", "DSP-C1", 99, 1,
            "A", "B", "C", "T", "S", "LR-PIECES", "2026-09-17",
        ])
        buffer = BytesIO()
        workbook.save(buffer)
        response = self.client.post(
            "/api/dispatches/import/preview/",
            {"file": SimpleUploadedFile("pieces.xlsx", buffer.getvalue())},
            format="multipart",
            HTTP_HOST="localhost",
        )
        self.assertEqual(response.status_code, 200, response.content)
        body = response.json()
        self.assertEqual(len(body["importable"]), 1)
        self.assertIn("matched by customer + item", body["importable"][0]["match_message"])
        token = body["preview_token"]
        commit = self.client.post(
            "/api/dispatches/import/commit/",
            {"preview_token": token, "group_keys": [body["importable"][0]["group_key"]]},
            format="json",
            HTTP_HOST="localhost",
        )
        self.assertEqual(commit.status_code, 201, commit.content)
        po.refresh_from_db()
        self.assertEqual(po.brand_order_code, "ATS-CUSTOMER-ITEM")
        self.assertEqual(DispatchLine.objects.get().dispatched_quantity, Decimal("1.00"))

    def test_customer_item_fallback_reports_ambiguity_and_missing_shade(self):
        customer = self.customer
        key_a = "TV-1-163-P01-S"
        key_b = key_a
        for index, key in enumerate((key_a, key_b), start=1):
            po = PurchaseOrder.objects.create(
                po_number=f"FALLBACK-{index}",
                entity=self.entity,
                customer=customer,
                brand=self.brand,
                po_date="2026-09-16",
                status=PurchaseOrder.Status.BRAND_ACCEPTED,
            )
            PurchaseOrderLine.objects.create(
                purchase_order=po,
                product_description=key,
                quantity=Decimal("5"),
                rate=Decimal("10"),
                line_total=Decimal("50"),
            )
        ambiguous = self._preview([_row("ATS-NONE", "INV-A", customer_code="DSP-C1", code1="TV", code5="1", code6="163", template="P01", size="S", qty=1)])
        self.assertEqual(ambiguous.status_code, 200, ambiguous.content)
        self.assertIn("FALLBACK-1", ambiguous.json()["excluded"][0]["reason"])
        self.assertIn("FALLBACK-2", ambiguous.json()["excluded"][0]["reason"])
        missing = self._preview([_row("ATS-NONE-2", "INV-B", customer_code="DSP-C1", code1="TV", code5="1", code6="999", template="P01", size="S", qty=1)])
        self.assertEqual(missing.status_code, 200, missing.content)
        self.assertIn("item keys missing from customer's open POs", missing.json()["excluded"][0]["reason"])

    def test_no_customer_item_match_reports_customer_master_without_pos(self):
        no_po_customer = Customer.objects.create(
            customer_code="NO-PO-CUSTOMER",
            customer_name="Customer Without POs",
        )
        CustomerEntity.objects.create(
            customer=no_po_customer,
            entity=self.entity,
            primary_entity=True,
        )
        response = self._preview(
            [_row("ATS-NO-PO", "INV-NO-PO", customer_code="NO-PO-CUSTOMER")]
        )
        self.assertEqual(response.status_code, 200, response.content)
        reason = response.json()["excluded"][0]["reason"]
        self.assertIn("customer in master: yes", reason)
        self.assertIn("POs for this customer: 0", reason)
        self.assertIn("item keys missing from customer's open POs", reason)

    def test_no_customer_item_match_lists_draft_po_status(self):
        self._create_po(
            "DRAFT-ONLY",
            status="Draft",
            description="A-B-C-T-S",
        )
        response = self._preview(
            [_row("ATS-DRAFT-ONLY", "INV-DRAFT", code1="A", code5="B", code6="C", template="T", size="S")]
        )
        self.assertEqual(response.status_code, 200, response.content)
        reason = response.json()["excluded"][0]["reason"]
        self.assertIn("customer in master: yes", reason)
        self.assertIn("POs for this customer: 1", reason)
        self.assertIn("DRAFT-ONLY is Draft; Brand Accepted or Partially Dispatched is required", reason)

    def test_no_customer_item_match_reports_missing_shade_key(self):
        self._create_po(
            "OPEN-WRONG-SHADE",
            status="Brand Accepted",
            description="A-B-OTHER-T-S",
        )
        response = self._preview(
            [_row("ATS-MISSING-SHADE", "INV-SHADE", code1="A", code5="B", code6="C", template="T", size="S")]
        )
        self.assertEqual(response.status_code, 200, response.content)
        reason = response.json()["excluded"][0]["reason"]
        self.assertIn("customer in master: yes", reason)
        self.assertIn("POs for this customer: 1", reason)
        self.assertIn("item keys missing from customer's open POs: A-B-C-T-S", reason)

    def test_fixture_imports_all_groups_and_lines(self):
        self._create_fixture_pos()
        preview = self.client.post(
            "/api/dispatches/import/preview/",
            {"file": self._fixture_upload()},
            format="multipart",
            HTTP_HOST="localhost",
        )
        self.assertEqual(preview.status_code, 200, preview.content)
        body = preview.json()
        self.assertEqual(len(body["importable"]), 12)
        self.assertEqual(sum(row["line_count"] for row in body["importable"]), 38)
        token = body["preview_token"]
        commit = self.client.post(
            "/api/dispatches/import/commit/",
            {"preview_token": token, "group_keys": [row["group_key"] for row in body["importable"]]},
            format="json",
            HTTP_HOST="localhost",
        )
        self.assertEqual(commit.status_code, 201, commit.content)
        self.assertEqual(commit.json()["created_count"], 12)
        self.assertEqual(
            Dispatch.objects.filter(challan_reference="DOM2601213").count(),
            2,
        )

    def test_fixture_wrong_customer_excludes_only_one_group(self):
        po_ids, customers = self._create_fixture_pos()
        wrong_customer = Customer.objects.create(
            customer_code="WRONG-CUSTOMER",
            customer_name="Wrong Customer",
        )
        po = PurchaseOrder.objects.get(pk=po_ids["ATS26RG0017"])
        po.customer = wrong_customer
        po.save(update_fields=["customer"])
        preview = self.client.post(
            "/api/dispatches/import/preview/",
            {"file": self._fixture_upload()},
            format="multipart",
            HTTP_HOST="localhost",
        )
        self.assertEqual(preview.status_code, 200, preview.content)
        body = preview.json()
        self.assertEqual(len(body["importable"]), 11)
        self.assertEqual(len(body["excluded"]), 1)
        reason = body["excluded"][0]["reason"]
        self.assertIn("700", reason)
        self.assertIn("WRONG-CUSTOMER", reason)

    def test_fixture_reimport_excludes_every_group(self):
        self._create_fixture_pos()
        preview = self.client.post(
            "/api/dispatches/import/preview/",
            {"file": self._fixture_upload()},
            format="multipart",
            HTTP_HOST="localhost",
        )
        self.assertEqual(preview.status_code, 200, preview.content)
        body = preview.json()
        commit = self.client.post(
            "/api/dispatches/import/commit/",
            {"preview_token": body["preview_token"], "group_keys": [row["group_key"] for row in body["importable"]]},
            format="json",
            HTTP_HOST="localhost",
        )
        self.assertEqual(commit.status_code, 201, commit.content)
        second = self.client.post(
            "/api/dispatches/import/preview/",
            {"file": self._fixture_upload()},
            format="multipart",
            HTTP_HOST="localhost",
        )
        self.assertEqual(second.status_code, 200, second.content)
        second_body = second.json()
        self.assertEqual(second_body["importable"], [])
        self.assertEqual(len(second_body["excluded"]), 12)
        self.assertTrue(all("already imported" in row["reason"] for row in second_body["excluded"]))
        self.assertEqual(Dispatch.objects.count(), 12)

    def test_missing_customer_cd_is_reported_with_other_required_columns(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Order Code"])
        sheet.append(["ATS-MISSING"])
        buffer = BytesIO()
        workbook.save(buffer)
        upload = SimpleUploadedFile("missing-customer.xlsx", buffer.getvalue())
        response = self.client.post(
            "/api/dispatches/import/preview/",
            {"file": upload},
            format="multipart",
            HTTP_HOST="localhost",
        )
        self.assertEqual(response.status_code, 400, response.content)
        message = str(response.json()["file"])
        self.assertIn("CustomerCd", message)
        self.assertIn("Invoice No.", message)
        self.assertIn("Quantity", message)

    def test_blank_brand_order_codes_are_allowed_but_duplicate_nonblank_is_friendly(self):
        first = self._create_po("BLANK-A")
        second = self._create_po("BLANK-B")
        self.assertNotEqual(first, second)
        first = self._create_po("DUP-A", brand_order_code="DUPLICATE")
        response = self.client.post(
            "/api/purchase-orders/",
            {
                "po_number": "DUP-B",
                "brand_order_code": "DUPLICATE",
                "entity_id": self.entity.id,
                "customer_id": self.customer.id,
                "brand_id": self.brand.id,
                "po_date": "2026-09-16",
                "lines": [
                    {
                        "product_id": None,
                        "product_description": "A-B-C-T-S",
                        "quantity": "10.00",
                        "rate": "10.00",
                    }
                ],
            },
            format="json",
            HTTP_HOST="localhost",
        )
        self.assertEqual(response.status_code, 400, response.content)
        self.assertIn("already used for this brand", str(response.json()))

    def test_product_code_matching_code1_is_a_fallback_without_client_item_code(self):
        product = Product.objects.create(
            product_code="TV",
            product_name="TV fallback",
            brand=self.brand,
            category=self.product.category,
            rate=Decimal("10.00"),
            unit=Product.Unit.METER,
        )
        po = PurchaseOrder.objects.create(
            po_number="FALLBACK-PO",
            brand_order_code="FALLBACK-ORDER",
            entity=self.entity,
            customer=self.customer,
            brand=self.brand,
            po_date="2026-09-16",
            status=PurchaseOrder.Status.BRAND_ACCEPTED,
        )
        PurchaseOrderLine.objects.create(
            purchase_order=po,
            product=product,
            quantity=Decimal("10.00"),
            rate=Decimal("10.00"),
            line_total=Decimal("10.00"),
        )
        response = self._preview([_row("FALLBACK-ORDER", "INV-FALLBACK", code1="TV")])
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(len(response.json()["importable"]), 1)

    def test_preview_no_matching_po_excludes_all(self):
        res = self._preview([_row("MISSING-PO", "INV-1")])
        self.assertEqual(res.status_code, 200, res.content)
        body = res.json()
        self.assertEqual(body["importable"], [])
        self.assertEqual(len(body["excluded"]), 1)
        self.assertIn("No customer + item PO match for Order Code MISSING-PO", body["excluded"][0]["reason"])
        self.assertEqual(Dispatch.objects.count(), 0)

    def test_preview_draft_po_excluded_not_brand_accepted(self):
        self._create_po("ATN-DRAFT", status="Draft")
        res = self._preview([_row("ATN-DRAFT", "INV-1")])
        self.assertEqual(res.status_code, 200, res.content)
        body = res.json()
        self.assertEqual(body["importable"], [])
        reason = body["excluded"][0]["reason"]
        self.assertIn("must be Brand Accepted or Partially Dispatched", reason)
        self.assertIn("Draft", reason)
        self.assertNotIn("No matching Purchase Order", reason)

    def test_preview_customer_mismatch_excluded(self):
        self._create_po("ATN-OK", status="Brand Accepted")
        res = self._preview([_row("ATN-OK", "INV-1", customer_code="DSP-C2", customer_name="Other House")])
        self.assertEqual(res.status_code, 200, res.content)
        body = res.json()
        self.assertEqual(body["importable"], [])
        reason = body["excluded"][0]["reason"]
        self.assertIn("Customer mismatch", reason)
        self.assertIn("DSP-C2", reason)
        self.assertIn("DSP-C1", reason)

    def test_preview_invoice_spanning_two_order_codes_splits(self):
        self._create_po("ATN-A", status="Brand Accepted")
        self._create_po("ATN-B", status="Brand Accepted")
        res = self._preview(
            [
                _row("ATN-A", "INV-SHARED", qty=3),
                _row("ATN-B", "INV-SHARED", qty=2),
            ]
        )
        self.assertEqual(res.status_code, 200, res.content)
        body = res.json()
        self.assertEqual(len(body["importable"]), 2)
        keys = {row["group_key"] for row in body["importable"]}
        self.assertEqual(keys, {"ATN-A||INV-SHARED", "ATN-B||INV-SHARED"})
        self.assertEqual(body["excluded"], [])

    def test_preview_over_dispatch_uses_pending_message(self):
        self._create_po("ATN-QTY", status="Brand Accepted", qty="4.00")
        res = self._preview([_row("ATN-QTY", "INV-1", qty=10)])
        self.assertEqual(res.status_code, 200, res.content)
        body = res.json()
        self.assertEqual(body["importable"], [])
        reason = body["excluded"][0]["reason"]
        self.assertIn("Cannot dispatch 10 units", reason)
        self.assertIn("only 4 units pending", reason)

    def test_commit_creates_dispatches_and_updates_po_status(self):
        po_id = self._create_po("ATN-FULL", status="Brand Accepted", qty="10.00")
        preview = self._preview(
            [
                _row("ATN-FULL", "INV-1", qty=4, lr_no="LR-A"),
                _row("ATN-FULL", "INV-2", qty=6, lr_no="LR-B"),
            ]
        )
        self.assertEqual(preview.status_code, 200, preview.content)
        token = preview.json()["preview_token"]
        keys = [row["group_key"] for row in preview.json()["importable"]]
        self.assertEqual(len(keys), 2)
        commit = self.client.post(
            "/api/dispatches/import/commit/",
            {"preview_token": token, "group_keys": keys},
            format="json",
            HTTP_HOST="localhost",
        )
        self.assertEqual(commit.status_code, 201, commit.content)
        body = commit.json()
        self.assertEqual(body["created_count"], 2)
        self.assertEqual(Dispatch.objects.count(), 2)
        po = PurchaseOrder.objects.get(pk=po_id)
        self.assertEqual(po.status, PurchaseOrder.Status.FULLY_DISPATCHED)
        statuses = {row["po_status"] for row in body["created"]}
        self.assertIn(PurchaseOrder.Status.FULLY_DISPATCHED, statuses)

    def test_commit_requires_preview_and_does_not_write_on_preview(self):
        self._create_po("ATN-PREV", status="Brand Accepted")
        preview = self._preview([_row("ATN-PREV", "INV-1", qty=2)])
        self.assertEqual(preview.status_code, 200, preview.content)
        self.assertEqual(Dispatch.objects.count(), 0)
        denied = self.client.post(
            "/api/dispatches/import/commit/",
            {"group_keys": ["ATN-PREV||INV-1"]},
            format="json",
            HTTP_HOST="localhost",
        )
        self.assertEqual(denied.status_code, 400)
        self.assertEqual(Dispatch.objects.count(), 0)

    def test_operator_can_preview_and_commit(self):
        self._create_po("ATN-OP", status="Brand Accepted", qty="5.00")
        PurchaseOrder.objects.filter(po_number="ATN-OP").update(created_by=self.operator)
        preview = self._preview([_row("ATN-OP", "INV-1", qty=5)], auth=self.operator)
        self.assertEqual(preview.status_code, 200, preview.content)
        token = preview.json()["preview_token"]
        keys = [row["group_key"] for row in preview.json()["importable"]]
        commit = self.client.post(
            "/api/dispatches/import/commit/",
            {"preview_token": token, "group_keys": keys},
            format="json",
            HTTP_HOST="localhost",
        )
        self.assertEqual(commit.status_code, 201, commit.content)
        po = PurchaseOrder.objects.get(po_number="ATN-OP")
        self.assertEqual(po.status, PurchaseOrder.Status.FULLY_DISPATCHED)
