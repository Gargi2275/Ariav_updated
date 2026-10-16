from __future__ import annotations

import hashlib
import re
import secrets
import zipfile
from collections import Counter, OrderedDict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.core.cache import cache
from django.db import transaction
from defusedxml import ElementTree as DefusedET
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

from brands.models import Brand
from entities.models import Entity

from .models import PriceList, PriceListImportBatch, PriceListQualityEntry, PriceListQualityShade
from .services import overlapping_active_warning

FORMAT_ID = "ocm_regular_pricelist"
CACHE_PREFIX = "price_list_quality_preview:"
TTL = 30 * 60
QUALITY_RE = re.compile(r"^[^-]+-[^-]+-[^-]+-[^-]+-[^-]+$")
TITLE_RE = re.compile(r"\b(Summer|Winter)\s*['’]?\s*(20\d{2})\b", re.I)
WEF_RE = re.compile(r"\bw\.?\s*e\.?\s*f\.?\s*[:.]?\s*(\d{1,2}[./-]\d{1,2}[./-]\d{4})", re.I)
PAGE_RE = re.compile(r"\bPage\s+(\d+)\s+of\s+(\d+)\b", re.I)
COLLECTION_RE = re.compile(r"^\s*(.+?)\s*\(\s*([^)]+)\s*\)\s*$")
MONTHS = {name.lower(): index for index, name in enumerate(
    ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"), 1
)}
CACHE_SIZE_NOTE = "Quality previews are cached for 30 minutes and retain all parsed rows; no row data is truncated."
MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_UNCOMPRESSED_SIZE = 50 * 1024 * 1024
MAX_PACKAGE_ENTRIES = 200
MAX_ROWS = 5000
MAX_COLUMNS = 50
REQUIRED_COLUMNS = ("serial", "quality", "shades", "delivery", "ex_mill", "wholesale", "retail")
COLUMN_LABELS = {
    "serial": "Serial No.",
    "quality": "Quality No.",
    "shades": "Shades",
    "delivery": "Delivery",
    "ex_mill": "Ex-Mill",
    "wholesale": "W/S",
    "retail": "Retail",
}


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header(value) -> str:
    return " ".join(_text(value).lower().replace(".", "").replace("_", " ").replace("-", " ").split())


def _validate_package(file_bytes):
    if len(file_bytes) > MAX_FILE_SIZE:
        raise ValidationError({"file": "The quality price list must be an .xlsx file no larger than 5 MB."})
    if not file_bytes.startswith(b"PK\x03\x04"):
        raise ValidationError({"file": "The uploaded file is not a valid .xlsx package."})
    try:
        with zipfile.ZipFile(BytesIO(file_bytes)) as package:
            entries = package.infolist()
            if len(entries) > MAX_PACKAGE_ENTRIES:
                raise ValidationError({"file": "The spreadsheet contains too many package entries."})
            if any(info.flag_bits & 0x1 for info in entries):
                raise ValidationError({"file": "Password-protected spreadsheets are not supported."})
            if sum(info.file_size for info in entries) > MAX_UNCOMPRESSED_SIZE:
                raise ValidationError({"file": "The spreadsheet expands beyond the 50 MB safety limit."})
            for info in entries:
                if info.filename.lower().endswith((".xml", ".rels")):
                    try:
                        DefusedET.fromstring(package.read(info))
                    except Exception as exc:
                        raise ValidationError({"file": "The spreadsheet contains corrupt or unsafe XML."}) from exc
                if "vbaproject" in info.filename.lower():
                    raise ValidationError({"file": "Macro-enabled spreadsheets are not supported."})
    except zipfile.BadZipFile as exc:
        raise ValidationError({"file": "The uploaded file is corrupt or not a valid .xlsx package."}) from exc


def _cell_matrix(sheet):
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) > MAX_ROWS:
        raise ValidationError({"file": "The spreadsheet contains more than 5,000 rows."})
    if max((len(row) for row in rows), default=0) > MAX_COLUMNS:
        raise ValidationError({"file": "The spreadsheet contains more than 50 columns."})
    return [tuple(row) for row in rows]


def _decimal(value):
    raw = _text(value).replace(",", "").replace("₹", "").replace(" ", "")
    if not raw or raw in {"-", "–"}:
        return None
    try:
        return Decimal(raw)
    except (InvalidOperation, ValueError):
        return None


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = _text(value)
    for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw[:10], fmt).date()
        except ValueError:
            pass
    return None


def _flag(severity, code, message):
    return {"severity": severity, "code": code, "message": message}


def _reason_summary(entries, severity):
    grouped = OrderedDict()
    for entry in entries:
        for flag in entry["flags"]:
            if flag["severity"] != severity:
                continue
            item = grouped.setdefault(flag["code"], {"code": flag["code"], "message": flag["message"], "count": 0, "examples": []})
            item["count"] += 1
            if len(item["examples"]) < 3:
                item["examples"].append(entry["source_row"])
    return list(grouped.values())


def _find_title(rows):
    title = ""
    season = year = None
    effective = None
    for row in rows[:40]:
        joined = " ".join(_text(value) for value in row if _text(value))
        match = TITLE_RE.search(joined)
        if match and not season:
            season, year = match.group(1).title(), int(match.group(2))
            title = joined
        if not effective:
            date_match = WEF_RE.search(joined)
            if date_match:
                effective = _date(date_match.group(1))
    if not season or not year:
        raise ValidationError({"file": "Could not detect Summer/Winter and a four-digit year in the first 40 rows."})
    if not title:
        title = f"{season} {year}"
    return title, season, year, effective


def _column_name(value):
    normalized = _header(value)
    if "serial" in normalized:
        return "serial"
    if "quality no" in normalized or normalized == "quality":
        return "quality"
    if "shade" in normalized:
        return "shades"
    if "delivery" in normalized:
        return "delivery"
    if "ex mill" in normalized or "exmill" in normalized:
        return "ex_mill"
    if normalized in {"w/s", "ws", "wholesale"}:
        return "wholesale"
    if "retail" in normalized:
        return "retail"
    return None


def _header_indexes(rows, sheet):
    quality_row = None
    for index, row in enumerate(rows[:40]):
        if any("quality no" in _header(value) or _header(value) == "quality" for value in row):
            quality_row = index
            break
    if quality_row is None:
        raise ValidationError({
            "code": "PL_HEADER_REQUIRED_MISSING",
            "file": "Required columns are missing. Found: none. Expected: Serial No., Quality No., Shades, Delivery, Ex-Mill, W/S, Retail. Header rows examined: 1-40.",
        })

    first_data = len(rows)
    for index in range(quality_row + 1, len(rows)):
        values = rows[index]
        joined = " ".join(_text(value) for value in values if _text(value))
        serial = next((value for value in values if _text(value).isdigit()), None)
        if serial is not None or COLLECTION_RE.match(joined):
            first_data = index
            break
    band_end = max(quality_row + 1, first_data)
    band_rows = list(range(quality_row, band_end))
    merged = {}
    for cell_range in sheet.merged_cells.ranges:
        min_col, min_row, max_col, max_row = cell_range.bounds
        if min_row - 1 <= quality_row < max_row and min_row - 1 < band_end:
            for row in range(max(quality_row, min_row - 1), min(band_end, max_row)):
                for column in range(min_col - 1, max_col):
                    merged[(row, column)] = (min_row - 1, min_col - 1)

    columns = {}
    unit_text = ""
    header_rows = [row + 1 for row in band_rows]
    effective_headers = []
    for column in range(max((len(rows[row]) for row in band_rows), default=0)):
        labels = []
        group_labels = []
        for row in band_rows:
            source = merged.get((row, column), (row, column))
            value = rows[source[0]][source[1]] if source[1] < len(rows[source[0]]) else None
            if _text(value) and _text(value) not in labels:
                labels.append(_text(value))
            if source[0] == row and _text(value) and source[1] != column:
                group_labels.append(_text(value))
        for label in labels:
            name = _column_name(label)
            if name and name not in columns:
                columns[name] = column
            elif not name and any(token in _header(label) for token in ("inr", "per mtr", "per meter")):
                unit_text = unit_text or label
        if group_labels:
            candidate = group_labels[0]
            if not _column_name(candidate):
                unit_text = unit_text or candidate
        effective_headers.append(labels[-1] if labels else "")
    if not unit_text:
        for cell_range in sheet.merged_cells.ranges:
            min_col, min_row, max_col, max_row = cell_range.bounds
            if min_row - 1 in band_rows and max_col > min_col:
                value = rows[min_row - 1][min_col - 1] if min_col - 1 < len(rows[min_row - 1]) else None
                if _text(value) and not _column_name(value):
                    unit_text = _text(value)
                    break
    missing = [name for name in REQUIRED_COLUMNS if name not in columns]
    if missing:
        found = [label for label in effective_headers if label]
        expected = ", ".join(COLUMN_LABELS[name] for name in REQUIRED_COLUMNS)
        examined = ", ".join(str(row) for row in header_rows)
        raise ValidationError({
            "code": "PL_HEADER_REQUIRED_MISSING",
            "file": f"Required columns are missing: {', '.join(COLUMN_LABELS[name] for name in missing)}. Found: {', '.join(found) or 'none'}. Expected: {expected}. Header rows examined: {examined}.",
            "missing_columns": [COLUMN_LABELS[name] for name in missing],
            "detected_headers": found,
            "header_rows": header_rows,
        })
    return quality_row, columns, effective_headers, unit_text, header_rows


def _split_shades(raw):
    return [token for token in re.split(r"[\s,\r\n]+", raw) if token]


def _shade_norm(token):
    value = token.strip().upper()
    if value.isdigit():
        return str(int(value))
    return value


def _parse_quality(value):
    raw = _text(value)
    parts = raw.split("-") if QUALITY_RE.match(raw) else []
    return raw, parts + [""] * (5 - len(parts))


def _months(raw):
    value = raw.casefold()
    for name, number in MONTHS.items():
        if name.startswith(value) and len(value) >= 3:
            return number
    return None


def _period(season, year, effective):
    start = effective or (date(year, 1, 1) if season == "Summer" else date(year, 7, 1))
    end = date(year, 6, 30) if season == "Summer" else date(year, 12, 31)
    return start, end


def parse_ocm(file_bytes: bytes, selected_year=None, selected_season=None):
    _validate_package(file_bytes)
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=False, data_only=True, keep_links=False, keep_vba=False)
    except Exception as exc:
        raise ValidationError({"file": f"Could not read the spreadsheet: {exc}"}) from exc
    try:
        sheet = workbook["Price List"] if "Price List" in workbook.sheetnames else workbook.active
        rows = _cell_matrix(sheet)
        if sheet.sheet_state != "visible":
            raise ValidationError({"file": "The Price List sheet is hidden and cannot be imported."})
        hidden_rows = [
            index for index, dimension in sheet.row_dimensions.items()
            if dimension.hidden and any(_text(value) for value in rows[index - 1])
        ]
        hidden_columns = [
            key for key, dimension in sheet.column_dimensions.items()
            if dimension.hidden and any(_text(row[dimension.min - 1 if dimension.min else 0]) for row in rows)
        ]
        header_row, columns, effective_headers, unit_text, header_rows = _header_indexes(rows, sheet)
        hidden_sheets = [
            other.title for other in workbook.worksheets
            if other.title != sheet.title and other.sheet_state != "visible"
            and any(_text(value) for row in other.iter_rows(values_only=True) for value in row)
        ]
        hidden_data = {"rows": hidden_rows, "columns": hidden_columns, "sheets": hidden_sheets}
    finally:
        workbook.close()
    title, detected_season, detected_year, effective = _find_title(rows)
    season = selected_season or detected_season
    year = int(selected_year or detected_year)
    if season not in ("Summer", "Winter") or not 2000 <= year <= 2100:
        raise ValidationError({"season": "Season must be Summer or Winter and year must be from 2000 to 2100."})
    start, end = _period(season, year, effective)
    header_values = tuple(effective_headers)
    known = set(columns.values())
    unrecognized = [
        {"header": _text(value), "non_empty_cells": sum(bool(_text(row[index] if index < len(row) else None)) for row in rows[header_row + 1:])}
        for index, value in enumerate(header_values) if index not in known and _text(value)
    ]
    counts = Counter()
    collections = OrderedDict()
    entries = []
    current_collection = {"name": "", "code": ""}
    current_page = None
    previous = None
    for row_number, values in enumerate(rows, 1):
        joined = " ".join(_text(value) for value in values if _text(value))
        page = PAGE_RE.search(joined)
        if page:
            current_page = int(page.group(1))
            counts["page"] += 1
            continue
        if row_number <= header_row:
            if joined:
                counts["title"] += 1
            else:
                counts["blank"] += 1
            continue
        if row_number in header_rows:
            counts["header"] += 1
            continue
        if not joined:
            counts["blank"] += 1
            continue
        quality_value = values[columns["quality"]] if columns["quality"] < len(values) else None
        quality = _text(quality_value)
        if _header(quality) in {"quality no", "quality"}:
            counts["header"] += 1
            continue
        serial_value = values[columns["serial"]] if "serial" in columns and columns["serial"] < len(values) else None
        price_values = {
            name: values[index] if index < len(values) else None
            for name, index in columns.items() if name in {"ex_mill", "wholesale", "retail"}
        }
        collection = COLLECTION_RE.match(quality) if quality else COLLECTION_RE.match(joined)
        if collection and not any(_text(value) for value in price_values.values()):
            current_collection = {"name": collection.group(1).strip(), "code": collection.group(2).strip()}
            collections[current_collection["code"]] = current_collection["name"]
            continue
        if not quality and previous and _text(values[columns["shades"]] if columns["shades"] < len(values) else None):
            continuation = _text(values[columns["shades"]])
            previous["shades_raw"] = f"{previous['shades_raw']}{continuation}"
            previous["shade_tokens"] = _split_shades(previous["shades_raw"])
            previous["raw_values"]["Shades"] = previous["shades_raw"]
            continue
        if not quality and not any(_text(value) for value in price_values.values()):
            collection = COLLECTION_RE.match(joined)
            if collection:
                current_collection = {"name": collection.group(1).strip(), "code": collection.group(2).strip()}
                collections[current_collection["code"]] = current_collection["name"]
            else:
                counts["header"] += 1
            continue
        if not quality:
            if previous and _text(values[columns["shades"]] if columns["shades"] < len(values) else None):
                continuation = _text(values[columns["shades"]])
                previous["shades_raw"] = f"{previous['shades_raw']}{continuation}"
                previous["raw_values"]["Shades"] = previous["shades_raw"]
            else:
                counts["blank"] += 1
            continue
        raw, codes = _parse_quality(quality)
        serial = None
        try:
            serial = int(_text(serial_value)) if _text(serial_value) else None
        except ValueError:
            pass
        shades_raw = _text(values[columns["shades"]] if columns["shades"] < len(values) else None)
        delivery_raw = _text(values[columns["delivery"]] if "delivery" in columns and columns["delivery"] < len(values) else None)
        flags = []
        if not QUALITY_RE.match(raw):
            flags.append(_flag("warning", "quality_pattern", "Quality No. does not contain five dash-separated segments."))
        if current_collection["code"] and codes[0] and codes[0].casefold() != current_collection["code"].casefold():
            flags.append(_flag("warning", "collection_prefix", "Quality No. prefix differs from the current collection code."))
        tokens = _split_shades(shades_raw)
        seen_tokens = set()
        for token in tokens:
            normalized = _shade_norm(token)
            if normalized in seen_tokens:
                flags.append(_flag("warning", "duplicate_shade", f"Shade {token} is repeated."))
            if re.match(r"^O\d", token, re.I):
                flags.append(_flag("warning", "possible_zero", f"Shade {token} may use letter O instead of zero."))
            if re.search(r"[^A-Za-z0-9]", token):
                flags.append(_flag("warning", "odd_shade", f"Shade {token} contains an unusual character."))
            seen_tokens.add(normalized)
        prices = {name: _decimal(price_values.get(name)) for name in ("ex_mill", "wholesale", "retail")}
        if prices["ex_mill"] is None:
            flags.append(_flag("error", "PL_PRICE_REQUIRED", "Ex-Mill price is required."))
        for name, value in prices.items():
            if value is not None and value <= 0:
                flags.append(_flag("error", "PL_PRICE_INVALID", f"{name.replace('_', ' ').title()} must be positive."))
        if prices["retail"] is None:
            flags.append(_flag("info", "missing_retail", "Retail price was not given."))
        if prices["wholesale"] is not None and prices["ex_mill"] is not None and prices["wholesale"] <= prices["ex_mill"]:
            flags.append(_flag("warning", "price_order", "W/S should be greater than Ex-Mill."))
        if prices["retail"] is not None and prices["wholesale"] is not None and prices["retail"] <= prices["wholesale"]:
            flags.append(_flag("warning", "price_order", "Retail should be greater than W/S."))
        delivery_month = _months(delivery_raw) if delivery_raw else None
        if delivery_raw and delivery_month is None:
            flags.append(_flag("warning", "delivery_month", f"Delivery '{delivery_raw}' is not a recognized month."))
        raw_values = {_text(header_values[i]) or f"Column {i + 1}": _text(values[i] if i < len(values) else None) for i in range(len(header_values))}
        item = {
            "source_row": row_number, "source_page": current_page, "collection_name": current_collection["name"],
            "collection_code": current_collection["code"], "serial_no": serial, "quality_no": raw,
            "code1": codes[0], "code2": codes[1], "code3": codes[2], "code4": codes[3], "code5": codes[4],
            "shades_raw": shades_raw, "shade_tokens": tokens, "delivery_raw": delivery_raw,
            "delivery_month": delivery_month, "ex_mill": prices["ex_mill"], "wholesale": prices["wholesale"],
            "retail": prices["retail"], "unit_text": unit_text, "flags": flags, "raw_values": raw_values,
        }
        entries.append(item)
        previous = item
        counts["row"] += 1
    by_collection = Counter(item["collection_code"] or item["collection_name"] for item in entries)
    for collection_code in by_collection:
        serials = [item["serial_no"] for item in entries if (item["collection_code"] or item["collection_name"]) == collection_code and item["serial_no"] is not None]
        expected = list(range(1, len(serials) + 1))
        if serials != expected:
            for item in entries:
                if (item["collection_code"] or item["collection_name"]) == collection_code:
                    item["flags"].append(_flag("warning", "serial_sequence", "Serial numbers must restart at 1 and have no gaps or repeats."))
    for quality, same in __import__("itertools").groupby(sorted(entries, key=lambda item: item["quality_no"]), key=lambda item: item["quality_no"]):
        rows_for_quality = list(same)
        if len(rows_for_quality) > 1:
            prices = {(row["ex_mill"], row["wholesale"], row["retail"]) for row in rows_for_quality}
            for row in rows_for_quality:
                row["flags"].append(_flag("error" if len(prices) > 1 else "warning", "duplicate_quality", f"Quality No. {quality} occurs more than once."))
    ratios = [item["wholesale"] / item["ex_mill"] for item in entries if item["ex_mill"] and item["wholesale"]]
    median = sorted(ratios)[len(ratios) // 2] if ratios else None
    if median:
        for item in entries:
            if item["ex_mill"] and item["wholesale"] and abs(item["wholesale"] / item["ex_mill"] - median) / median > Decimal("0.05"):
                item["flags"].append(_flag("warning", "ratio_outlier", "W/S differs by more than 5% from the list median ratio."))
    declared = max((int(match.group(2)) for match in (PAGE_RE.search(" ".join(_text(value) for value in row if _text(value))) for row in rows) if match), default=None)
    return {
        "format_id": FORMAT_ID, "title": title, "season": season, "year": year, "valid_from": start.isoformat(),
        "valid_to": end.isoformat(), "detected_effective_date": effective.isoformat() if effective else None,
        "entries": entries, "rows_read": counts["row"], "skipped_lines": {"blank": counts["blank"], "header": counts["header"], "page": counts["page"], "title": counts["title"]},
        "collections": [{"code": code, "name": name, "count": by_collection[code]} for code, name in collections.items()],
        "unrecognized_columns": unrecognized, "declared_page_count": declared, "cache_note": CACHE_SIZE_NOTE,
        "unit_text": unit_text,
        "header_rows": header_rows, "detected_columns": [
            {"column": index + 1, "header": effective_headers[index], "field": next((name for name, value in columns.items() if value == index), None)}
            for index in range(len(effective_headers)) if effective_headers[index]
        ], "hidden_data": hidden_data,
    }


def _validate_scope(branch_id, brand_id):
    branch = Entity.objects.filter(pk=branch_id, status=Entity.Status.ACTIVE).first()
    brand = Brand.objects.filter(pk=brand_id, status=Brand.Status.ACTIVE).first()
    if branch is None:
        raise ValidationError({"branch_id": "Select an active Branch."})
    if brand is None or not branch.brands.filter(pk=brand.pk).exists():
        raise ValidationError({"brand_id": "The selected Brand is not mapped to the selected Branch."})
    return branch, brand


def _public_preview(preview, token):
    entries = []
    for item in preview["entries"]:
        value = dict(item)
        for name in ("ex_mill", "wholesale", "retail"):
            if value[name] is not None:
                value[name] = str(value[name])
        entries.append(value)
    result = {key: value for key, value in preview.items() if key != "entries"}
    result["preview_token"] = token
    result["entries"] = entries
    global_warning_count = sum(item.get("count", 1) for item in result.get("warnings", []) if isinstance(item, dict))
    result["flagged_counts"] = {
        "error": sum(1 for item in entries for flag in item["flags"] if flag["severity"] == "error"),
        "warning": sum(1 for item in entries for flag in item["flags"] if flag["severity"] == "warning") + global_warning_count,
        "info": sum(1 for item in entries for flag in item["flags"] if flag["severity"] == "info"),
    }
    return result


def build_quality_preview(file_bytes, *, file_name, branch_id, brand_id, year=None, season=None, valid_from=None, valid_to=None):
    branch, brand = _validate_scope(branch_id, brand_id)
    digest = hashlib.sha256(file_bytes).hexdigest()
    existing = PriceListImportBatch.objects.filter(brand=brand, file_sha256=digest).select_related("price_list").first()
    if existing:
        raise ValidationError({"file": f"This file was already imported for {brand.brand_name} as {existing.price_list.season_label}."})
    preview = parse_ocm(file_bytes, year, season)
    if valid_from:
        preview["valid_from"] = valid_from
    if valid_to:
        preview["valid_to"] = valid_to
    preview.update({"branch_id": branch.id, "brand_id": brand.id, "file_name": file_name, "file_sha256": digest, "file_size": len(file_bytes), "warnings": [
        {"code": "PL_COLUMN_UNMAPPED", "message": f"Unrecognized column '{item['header']}' has {item['non_empty_cells']} non-empty cells.", "count": 1, "examples": []}
        for item in preview["unrecognized_columns"]
    ]})
    for kind, code, message in (
        ("rows", "PL_HIDDEN_ROWS", "Hidden rows contain data."),
        ("columns", "PL_HIDDEN_COLUMNS", "Hidden columns contain data."),
        ("sheets", "PL_HIDDEN_SHEETS", "Hidden sheets contain data."),
    ):
        if preview["hidden_data"].get(kind):
            preview["warnings"].append({"code": code, "message": f"{message} ({', '.join(map(str, preview['hidden_data'][kind]))}).", "count": len(preview["hidden_data"][kind]), "examples": preview["hidden_data"][kind][:3]})
    active = PriceList.objects.filter(brand=brand, status=PriceList.Status.ACTIVE, valid_from__lte=preview["valid_to"], valid_to__gte=preview["valid_from"]).first()
    if active:
        preview["warnings"].append({"code": "PL_ACTIVE_OVERLAP", "message": f"This {brand.brand_name} import ({preview['title']}, {preview['valid_from']} to {preview['valid_to']}) overlaps the Active list {active.season_label} ({active.valid_from} to {active.valid_to}, status Active). Importing creates a Draft and does not change the Active list until it is activated.", "count": 1, "examples": []})
    preview["warning_summary"] = _reason_summary(preview["entries"], "warning")
    preview["error_summary"] = _reason_summary(preview["entries"], "error")
    preview["info_summary"] = _reason_summary(preview["entries"], "info")
    preview["reconciliation"] = {
        "rows_read": preview["rows_read"] + sum(preview["skipped_lines"].values()),
        "imported": len(preview["entries"]),
        "excluded": 0,
        "skipped": sum(preview["skipped_lines"].values()),
        "skipped_by_kind": preview["skipped_lines"],
    }
    return preview


def store_quality_preview(user_id, preview):
    token = secrets.token_urlsafe(32)
    cache.set(f"{CACHE_PREFIX}{token}", {"user_id": user_id, **preview}, timeout=TTL)
    return token


def load_quality_preview(token):
    value = cache.get(f"{CACHE_PREFIX}{token}") if token else None
    return value if isinstance(value, dict) else None


def commit_quality_import(*, user, token, branch_id, brand_id, confirmed_warnings=False):
    preview = load_quality_preview(token)
    if not preview or preview.get("user_id") != getattr(user, "id", None):
        raise ValidationError({"detail": "Quality preview has expired or was not found. Run preview again."})
    _validate_scope(branch_id, brand_id)
    errors = [flag["message"] for item in preview["entries"] for flag in item["flags"] if flag["severity"] == "error"]
    warnings = [item["message"] if isinstance(item, dict) else item for item in preview.get("warnings", [])]
    warnings.extend(flag["message"] for item in preview["entries"] for flag in item["flags"] if flag["severity"] == "warning")
    if errors:
        raise ValidationError({"detail": "Resolve all error flags before committing.", "errors": errors})
    if warnings and not confirmed_warnings:
        raise ValidationError({"detail": "Confirm that you have reviewed the warnings before committing.", "warnings": warnings})
    brand = Brand.objects.get(pk=brand_id)
    if PriceListImportBatch.objects.filter(brand=brand, file_sha256=preview["file_sha256"]).exists():
        raise ValidationError({"file": f"This file was already imported for {brand.brand_name}."})
    with transaction.atomic():
        price_list = PriceList.objects.create(
            brand=brand, season_label=preview["title"], valid_from=date.fromisoformat(preview["valid_from"]),
            valid_to=date.fromisoformat(preview["valid_to"]), status=PriceList.Status.DRAFT, created_by=user,
        )
        batch = PriceListImportBatch.objects.create(
            price_list=price_list, brand=brand, format_id=FORMAT_ID, file_name=preview["file_name"],
            file_sha256=preview["file_sha256"], file_size=preview["file_size"], imported_by=user,
            rows_read=preview["rows_read"], rows_imported=len(preview["entries"]),
            skipped_blank=preview["skipped_lines"]["blank"], skipped_header=preview["skipped_lines"]["header"],
            skipped_page=preview["skipped_lines"]["page"], skipped_title=preview["skipped_lines"]["title"],
            collection_count=len(preview["collections"]), warnings=preview["warnings"],
            declared_page_count=preview["declared_page_count"],
        )
        for item in preview["entries"]:
            entry = PriceListQualityEntry.objects.create(
                price_list=price_list, batch=batch, **{key: item[key] for key in (
                    "source_row", "source_page", "collection_name", "collection_code", "serial_no", "quality_no",
                    "code1", "code2", "code3", "code4", "code5", "shades_raw", "delivery_raw", "delivery_month",
                    "ex_mill", "wholesale", "retail", "unit_text", "flags", "raw_values",
                )}
            )
            PriceListQualityShade.objects.bulk_create([
                PriceListQualityShade(entry=entry, position=index, shade_raw=token, shade_norm=_shade_norm(token))
                for index, token in enumerate(item["shade_tokens"], 1)
            ])
    cache.delete(f"{CACHE_PREFIX}{token}")
    return {"price_list_id": price_list.id, "batch_id": batch.id, "rows_imported": len(preview["entries"]), "season_label": price_list.season_label, "status": price_list.status, "warning_count": len(warnings)}
