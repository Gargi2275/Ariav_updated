from datetime import date
from decimal import Decimal
from io import BytesIO

from django.core.cache import cache
from django.test import TestCase
from openpyxl import Workbook
from rest_framework.test import APIClient

from accounts.models import AuthUser
from brands.models import Brand
from categories.models import Category
from products.models import Product
from entities.models import Entity

from .models import PriceList, PriceListEntry


class PriceListImportApiTests(TestCase):
    def setUp(self):
        cache.clear()
        self.user = AuthUser.objects.create_user(username="price-import-admin", password="pass12345", role=AuthUser.Role.ADMIN)
        self.brand = Brand.objects.create(brand_code="BR-1", brand_name="Brand One", order_method="Digital PO")
        self.other_brand = Brand.objects.create(brand_code="BR-2", brand_name="Brand Two", order_method="Digital PO")
        parent = Category.objects.create(category_code="CAT", category_name="Fabric")
        category = Category.objects.create(category_code="SUB", category_name="Cotton", parent_category=parent)
        self.product = Product.objects.create(product_code="SKU-1", product_name="Blue", brand=self.brand, category=category, rate=Decimal("40"))
        self.other_product = Product.objects.create(product_code="SKU-2", product_name="Red", brand=self.other_brand, category=category, rate=Decimal("41"))
        self.branch = Entity.objects.create(short_code="BR", entity_name="Test Branch", entity_type="Branch", status="Active")
        self.branch.brands.add(self.brand)
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def workbook(self, rows):
        book = Workbook()
        sheet = book.active
        sheet.title = "Price List"
        sheet.append(["Brand Code", "Product Code", "Product Name", "Taka Length (meters)", "Price Basis", "Price", "Taka Qty", "Season Label", "Valid From", "Valid To"])
        for row in rows:
            sheet.append(row)
        output = BytesIO()
        book.save(output)
        return output.getvalue()

    def workbook_without_season_columns(self, rows):
        book = Workbook()
        sheet = book.active
        sheet.title = "Price List"
        sheet.append(["Brand Code", "Product Code", "Product Name", "Taka Length (meters)", "Price Basis", "Price", "Taka Qty"])
        for row in rows:
            sheet.append(row[:7])
        output = BytesIO()
        book.save(output)
        return output.getvalue()

    def preview(self, rows, year=2026, season="Winter"):
        return self.client.post("/api/price-lists/import/preview/", {"file": BytesIO(self.workbook(rows)), "year": year, "season": season, "branch_id": self.branch.id, "brand_id": self.brand.id}, format="multipart")

    def valid_row(self, **overrides):
        row = ["BR-1", "SKU-1", "Reference", 25, "Per Taka", 1250, 10, "Winter 2026", date(2026, 7, 1), date(2026, 12, 31)]
        fields = ["brand", "product", "name", "length", "basis", "price", "qty", "season", "from", "to"]
        data = dict(zip(fields, row))
        data.update(overrides)
        return [data[key] for key in fields]

    def test_preview_matches_and_calculates_without_writing(self):
        response = self.preview([self.valid_row()])
        self.assertEqual(response.status_code, 200, response.content)
        body = response.json()
        self.assertEqual(len(body["importable"]), 1)
        self.assertEqual(body["importable"][0]["entries"][0]["price_per_meter"], "50.00")
        self.assertEqual(body["excluded"], [])
        self.assertEqual(PriceList.objects.count(), 0)

    def test_missing_column_error_reports_headers_and_normalizes_matching(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Price List"
        sheet.append([" Brand_Code ", "PRODUCT CODE", "Product Name", "Taka Length (meters)", "Price Basis", " Price_Value ", "Season Label", "Valid From", "Valid To"])
        output = BytesIO()
        workbook.save(output)
        response = self.client.post("/api/price-lists/import/preview/", {"file": BytesIO(output.getvalue()), "year": 2026, "season": "Winter", "branch_id": self.branch.id, "brand_id": self.brand.id}, format="multipart")
        self.assertEqual(response.status_code, 200, response.content)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Price List"
        sheet.append(["brand_code", "product_code"])
        output = BytesIO()
        workbook.save(output)
        response = self.client.post("/api/price-lists/import/preview/", {"file": BytesIO(output.getvalue()), "year": 2026, "season": "Winter", "branch_id": self.branch.id, "brand_id": self.brand.id}, format="multipart")
        self.assertEqual(response.status_code, 400, response.content)
        self.assertEqual(response.json()["detected_headers"], ["brand_code", "product_code"])
        self.assertIn("price_value", response.json()["file"])

    def test_preview_specific_exclusions_and_duplicate(self):
        rows = [
            self.valid_row(),
            self.valid_row(),
            self.valid_row(brand="UNKNOWN"),
            self.valid_row(product="SKU-2"),
            self.valid_row(basis="Per Yard"),
        ]
        response = self.preview(rows)
        self.assertEqual(response.status_code, 200, response.content)
        reasons = [row["reason"] for row in response.json()["excluded"]]
        self.assertIn("Duplicate product code within this price list", reasons)
        self.assertIn("Brand code UNKNOWN not found", reasons)
        self.assertTrue(any("exists under Brand Brand Two, not Brand Brand One" in reason for reason in reasons))
        self.assertTrue(any("Invalid Price Basis" in reason for reason in reasons))

    def test_commit_creates_active_and_rerun_excludes_existing(self):
        preview = self.preview([self.valid_row()]).json()
        commit = self.client.post("/api/price-lists/import/commit/", {"preview_token": preview["preview_token"], "group_keys": [preview["importable"][0]["import_key"]], "year": 2026, "season": "Winter", "branch_id": self.branch.id, "brand_id": self.brand.id}, format="json")
        self.assertEqual(commit.status_code, 201, commit.content)
        self.assertEqual(PriceList.objects.count(), 1)
        self.assertEqual(PriceList.objects.first().status, PriceList.Status.ACTIVE)
        self.assertEqual(PriceListEntry.objects.count(), 1)
        rerun = self.preview([self.valid_row()]).json()
        self.assertEqual(rerun["importable"], [])
        self.assertIn("A Winter 2026 list for Brand One already exists (Active)", rerun["excluded"][0]["reason"])

    def test_selected_season_ranges_and_optional_file_columns(self):
        summer = self.preview([self.valid_row(season="", **{"from": None, "to": None})], season="Summer").json()
        self.assertEqual(summer["valid_from"], "2026-01-01")
        self.assertEqual(summer["valid_to"], "2026-06-30")
        self.assertEqual(summer["season_label"], "Summer 2026")

        response = self.client.post(
            "/api/price-lists/import/preview/",
            {"file": BytesIO(self.workbook_without_season_columns([self.valid_row()])), "year": 2026, "season": "Winter", "branch_id": self.branch.id, "brand_id": self.brand.id},
            format="multipart",
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["importable"][0]["season_label"], "Winter 2026")

        leap = self.preview([self.valid_row(season="", **{"from": None, "to": None})], year=2028, season="Summer").json()
        self.assertEqual((leap["valid_from"], leap["valid_to"]), ("2028-01-01", "2028-06-30"))

    def test_file_dates_must_match_selected_season(self):
        rows = [
            self.valid_row(**{"from": date(2026, 1, 1), "to": date(2026, 6, 30)}),
            self.valid_row(product="SKU-1", **{"from": date(2026, 7, 1), "to": date(2026, 12, 31)}),
        ]
        response = self.preview(rows)
        self.assertEqual(response.status_code, 200, response.content)
        body = response.json()
        self.assertEqual(len(body["importable"]), 1)
        self.assertTrue(any("File says 2026-01-01 to 2026-06-30, you selected Winter 2026" in row["reason"] for row in body["excluded"]))

    def test_missing_or_invalid_selection_returns_400(self):
        payload = {"file": BytesIO(self.workbook([self.valid_row()])), "branch_id": self.branch.id, "brand_id": self.brand.id}
        missing = self.client.post("/api/price-lists/import/preview/", payload, format="multipart")
        self.assertEqual(missing.status_code, 400)
        invalid = self.client.post("/api/price-lists/import/preview/", {**payload, "year": 1999, "season": "Monsoon"}, format="multipart")
        self.assertEqual(invalid.status_code, 400)

    def test_existing_price_list_is_warned_and_not_duplicated(self):
        PriceList.objects.create(brand=self.brand, season_label="Winter 2026", valid_from=date(2026, 7, 1), valid_to=date(2026, 12, 31), status=PriceList.Status.DRAFT)
        response = self.preview([self.valid_row()])
        self.assertEqual(response.status_code, 200, response.content)
        body = response.json()
        self.assertEqual(body["importable"], [])
        self.assertIn("A Winter 2026 list for Brand One already exists (Draft)", body["warnings"])

    def test_commit_rejects_changed_stored_selection(self):
        preview = self.preview([self.valid_row()]).json()
        response = self.client.post(
            "/api/price-lists/import/commit/",
            {"preview_token": preview["preview_token"], "group_keys": [preview["importable"][0]["import_key"]], "year": 2025, "season": "Winter", "branch_id": self.branch.id, "brand_id": self.brand.id},
            format="json",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("do not match", response.json()["detail"])
