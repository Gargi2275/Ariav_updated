from __future__ import annotations

import hashlib
import secrets
from collections import OrderedDict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.core.cache import cache
from django.db import transaction
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

from brands.models import Brand
from entities.models import Entity
from products.models import Product

from .models import PriceList, PriceListEntry

PREVIEW_CACHE_PREFIX = "price_list_import_preview:"
PREVIEW_TTL_SECONDS = 30 * 60

HEADER_ALIASES = {
    "brand code": "brand_code",
    "brand": "brand_code",
    "product code": "product_code",
    "product name": "product_name",
    "taka length meters": "taka_length_meters",
    "taka length (meters)": "taka_length_meters",
    "taka length": "taka_length_meters",
    "price basis": "price_basis",
    "price": "price_value",
    "price value": "price_value",
    "taka qty": "taka_quantity",
    "taka quantity": "taka_quantity",
    "season label": "season_label",
    "valid from": "valid_from",
    "valid to": "valid_to",
}


def season_details(year: int, season: str) -> dict:
    if not isinstance(year, int) or isinstance(year, bool) or not 2000 <= year <= 2100:
        raise ValidationError({"year": "Year must be an integer from 2000 to 2100."})
    if season not in ("Summer", "Winter"):
        raise ValidationError({"season": "Season must be Summer or Winter."})
    if season == "Summer":
        valid_from, valid_to = date(year, 1, 1), date(year, 6, 30)
    else:
        valid_from, valid_to = date(year, 7, 1), date(year, 12, 31)
    return {
        "year": year,
        "season": season,
        "label": f"{season} {year}",
        "valid_from": valid_from,
        "valid_to": valid_to,
    }


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _norm_header(value) -> str:
    return " ".join(_cell_str(value).lower().replace("_", " ").replace(".", " ").split())


def _display_header(value) -> str:
    if value is None:
        return ""
    return str(value)


def _parse_decimal(value) -> Decimal | None:
    raw = _cell_str(value).replace(",", "")
    if not raw:
        return None
    try:
        return Decimal(raw)
    except (InvalidOperation, ValueError):
        return None


def _parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
        except (OverflowError, ValueError):
            return None
    raw = _cell_str(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(raw[:10], fmt).date()
        except ValueError:
            continue
    return None


def _date_text(value) -> str:
    parsed = _parse_date(value)
    return parsed.isoformat() if parsed else _cell_str(value)


def parse_workbook(file_bytes: bytes) -> tuple[list[dict], list[str]]:
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError({"file": f"Could not read the spreadsheet: {exc}"}) from exc

    warnings: list[str] = []
    try:
        sheet = workbook["Price List"] if "Price List" in workbook.sheetnames else workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            raise ValidationError({"file": "The spreadsheet is empty."}) from None
        detected_headers = [_display_header(cell) for cell in (header or ())]
        mapping = {index: HEADER_ALIASES[_norm_header(cell)] for index, cell in enumerate(header or ()) if _norm_header(cell) in HEADER_ALIASES}
        required = {"brand_code", "product_code", "taka_length_meters", "price_basis", "price_value"}
        missing = sorted(required - set(mapping.values()))
        if missing:
            raise ValidationError({
                "file": f"Missing required column(s): {', '.join(missing)}.",
                "detected_headers": detected_headers,
            })

        records = []
        for excel_row, values in enumerate(rows, start=2):
            record = {field: values[index] if index < len(values) else None for index, field in mapping.items()}
            brand_code = _cell_str(record.get("brand_code"))
            if not brand_code:
                break
            record["excel_row"] = excel_row
            record["brand_code"] = brand_code
            record["product_code"] = _cell_str(record.get("product_code"))
            record["product_name"] = _cell_str(record.get("product_name"))
            record["season_label"] = _cell_str(record.get("season_label"))
            record["price_basis"] = _cell_str(record.get("price_basis"))
            record["valid_from"] = _date_text(record.get("valid_from"))
            record["valid_to"] = _date_text(record.get("valid_to"))
            records.append(record)
    finally:
        workbook.close()
    return records, warnings


def _brand_map() -> dict[str, Brand]:
    return {(brand.brand_code or "").strip().lower(): brand for brand in Brand.objects.all()}


def _product_by_code(code: str) -> list[Product]:
    return list(Product.objects.select_related("brand").filter(product_code__iexact=code))


def _group_key(record: dict) -> str:
    return "||".join((record["brand_code"].lower(), record["season_label"], record["valid_from"], record["valid_to"]))


def _season_row_reason(record: dict, details: dict) -> str | None:
    selected_label = details["label"]
    file_label = record.get("season_label", "")
    if file_label and file_label.casefold() != selected_label.casefold():
        return f"File says season {file_label}, you selected {selected_label}"

    selected_from = details["valid_from"]
    selected_to = details["valid_to"]
    file_from_raw = record.get("valid_from", "")
    file_to_raw = record.get("valid_to", "")
    file_from = _parse_date(file_from_raw) if file_from_raw else None
    file_to = _parse_date(file_to_raw) if file_to_raw else None
    if file_from_raw and file_from is None:
        return f"File date '{_cell_str(file_from_raw)}' is invalid; you selected {selected_label}"
    if file_to_raw and file_to is None:
        return f"File date '{_cell_str(file_to_raw)}' is invalid; you selected {selected_label}"
    if (file_from and file_from != selected_from) or (file_to and file_to != selected_to):
        shown_from = file_from.isoformat() if file_from else selected_from.isoformat()
        shown_to = file_to.isoformat() if file_to else selected_to.isoformat()
        return f"File says {shown_from} to {shown_to}, you selected {selected_label}"
    return None


def _row_reason(record: dict, brands: dict[str, Brand], seen: set[tuple[str, str]]) -> tuple[Brand | None, Product | None, str | None]:
    brand = brands.get(record["brand_code"].lower())
    if brand is None:
        return None, None, f"Brand code {record['brand_code']} not found"
    products = _product_by_code(record["product_code"])
    product = next((item for item in products if item.brand_id == brand.id), None)
    if product is None:
        if products:
            other = products[0].brand.brand_name
            return brand, None, f"Product code {record['product_code']} exists under Brand {other}, not Brand {brand.brand_name}"
        return brand, None, f"Product code {record['product_code']} not found under Brand {brand.brand_name}"
    basis = record["price_basis"].lower()
    if basis not in ("per meter", "per taka"):
        return brand, product, f"Invalid Price Basis '{record['price_basis']}'. Expected Per Meter or Per Taka."
    length = _parse_decimal(record.get("taka_length_meters"))
    if length is None or length <= 0:
        return brand, product, "Taka Length must be a positive number"
    price = _parse_decimal(record.get("price_value"))
    if price is None or price <= 0:
        return brand, product, "Price must be a positive number"
    valid_from = _parse_date(record["valid_from"])
    valid_to = _parse_date(record["valid_to"])
    if valid_from is None or valid_to is None:
        return brand, product, "Valid From and Valid To must be valid dates"
    if valid_to <= valid_from:
        return brand, product, "Valid To must be after Valid From"
    duplicate_key = (_group_key(record), record["product_code"].lower())
    if duplicate_key in seen:
        return brand, product, "Duplicate product code within this price list"
    existing = PriceListEntry.objects.filter(
        product=product,
        price_list__brand=brand,
        price_list__season_label=record["season_label"],
        price_list__valid_from=valid_from,
        price_list__valid_to=valid_to,
    ).exists()
    if existing:
        return brand, product, "Already exists as an active price list entry"
    return brand, product, None


def _public_entry(record: dict, product: Product) -> dict:
    length = _parse_decimal(record["taka_length_meters"])
    price = _parse_decimal(record["price_value"])
    basis = "Per Taka" if record["price_basis"].lower() == "per taka" else "Per Meter"
    draft = PriceListEntry(taka_length_meters=length, price_basis=basis, price_value=price)
    draft.calculate_prices()
    return {
        "source_row": record["excel_row"],
        "product_code": product.product_code,
        "product_name": product.product_name,
        "product_id": product.id,
        "taka_length_meters": str(length),
        "price_basis": basis,
        "price_value": str(price),
        "price_per_meter": str(draft.price_per_meter),
        "price_per_taka": str(draft.price_per_taka),
        "taka_quantity": _cell_str(record.get("taka_quantity")) or None,
        "status": "Active",
    }


def build_preview(file_bytes: bytes, year: int, season: str, branch_id: int, brand_id: int) -> dict:
    details = season_details(year, season)
    branch = Entity.objects.filter(pk=branch_id, status=Entity.Status.ACTIVE).first()
    if branch is None:
        raise ValidationError({"branch_id": "Select an active Branch."})
    selected_brand = Brand.objects.filter(pk=brand_id, status=Brand.Status.ACTIVE).first()
    if selected_brand is None:
        raise ValidationError({"brand_id": "Select an active Brand."})
    if not branch.brands.filter(pk=selected_brand.pk).exists():
        raise ValidationError({"brand_id": "The selected Brand is not mapped to the selected Branch."})
    records, warnings = parse_workbook(file_bytes)
    brands = _brand_map()
    groups: OrderedDict[str, dict] = OrderedDict()
    excluded = []
    seen: set[tuple[str, str]] = set()
    duplicate_warned: set[int] = set()
    overlap_warned: set[int] = set()
    for record in records:
        record_brand = brands.get(record["brand_code"].lower())
        if record_brand is not None and record_brand.id != selected_brand.id:
            excluded.append({"source_row": record["excel_row"], "brand_code": record["brand_code"], "product_code": record["product_code"], "reason": f"Brand {record['brand_code']} is not the selected Brand for this Branch"})
            continue
        brand = brands.get(record["brand_code"].lower())
        season_reason = _season_row_reason(record, details)
        record["season_label"] = details["label"]
        record["valid_from"] = details["valid_from"].isoformat()
        record["valid_to"] = details["valid_to"].isoformat()
        brand, product, reason = _row_reason(record, brands, seen)
        if season_reason:
            reason = season_reason
        key = _group_key(record)
        group = groups.setdefault(key, {"import_key": key, "brand_code": record["brand_code"], "season_label": details["label"], "valid_from": details["valid_from"].isoformat(), "valid_to": details["valid_to"].isoformat(), "brand_id": None, "brand_name": record["brand_code"], "entries": []})
        if brand:
            group["brand_id"] = brand.id
            group["brand_name"] = brand.brand_name
            existing = PriceList.objects.filter(
                brand=brand,
                season_label=details["label"],
                valid_from=details["valid_from"],
                valid_to=details["valid_to"],
            ).first()
            if existing:
                reason = f"A {details['label']} list for {brand.brand_name} already exists ({existing.status})"
                if brand.id not in duplicate_warned:
                    warnings.append(reason)
                    duplicate_warned.add(brand.id)
            elif brand.id not in overlap_warned and PriceList.objects.filter(
                brand=brand,
                status=PriceList.Status.ACTIVE,
                valid_from__lte=details["valid_to"],
                valid_to__gte=details["valid_from"],
            ).exists():
                warnings.append(f"This {details['label']} import overlaps an Active list for {brand.brand_name}.")
                overlap_warned.add(brand.id)
        if reason:
            excluded.append({"source_row": record["excel_row"], "brand_code": record["brand_code"], "product_code": record["product_code"], "reason": reason})
            continue
        seen.add((key, record["product_code"].lower()))
        group["entries"].append(_public_entry(record, product))
    importable = [group for group in groups.values() if group["entries"]]
    return {
        "file_sha256": hashlib.sha256(file_bytes).hexdigest(),
        "year": details["year"],
        "season": details["season"],
        "season_label": details["label"],
        "valid_from": details["valid_from"].isoformat(),
        "valid_to": details["valid_to"].isoformat(),
        "branch_id": branch.id,
        "brand_id": selected_brand.id,
        "importable": importable,
        "excluded": excluded,
        "warnings": warnings,
    }


def store_preview(user_id: int, preview: dict) -> str:
    token = secrets.token_urlsafe(32)
    cache.set(f"{PREVIEW_CACHE_PREFIX}{token}", {"user_id": user_id, **preview}, timeout=PREVIEW_TTL_SECONDS)
    return token


def load_preview(token: str) -> dict | None:
    cached = cache.get(f"{PREVIEW_CACHE_PREFIX}{token}") if token else None
    return cached if isinstance(cached, dict) else None


def public_preview_payload(preview: dict, token: str) -> dict:
    return {
        "preview_token": token,
        "year": preview["year"],
        "season": preview["season"],
        "season_label": preview["season_label"],
        "valid_from": preview["valid_from"],
        "valid_to": preview["valid_to"],
        "importable": preview["importable"],
        "excluded": preview["excluded"],
        "warnings": preview["warnings"],
    }


def commit_import(*, user, token: str, group_keys: list[str], year: int, season: str, branch_id: int, brand_id: int) -> dict:
    cached = load_preview(token)
    if not cached or cached.get("user_id") != getattr(user, "id", None):
        raise ValidationError({"detail": "Preview has expired or was not found. Run preview again before committing."})
    details = season_details(year, season)
    branch = Entity.objects.filter(pk=branch_id, status=Entity.Status.ACTIVE).first()
    selected_brand = Brand.objects.filter(pk=brand_id, status=Brand.Status.ACTIVE).first()
    if branch is None:
        raise ValidationError({"branch_id": "Select an active Branch."})
    if selected_brand is None or not branch.brands.filter(pk=selected_brand.pk).exists():
        raise ValidationError({"brand_id": "The selected Brand is not mapped to the selected Branch."})
    if (
        cached.get("year") != details["year"]
        or cached.get("season") != details["season"]
        or cached.get("season_label") != details["label"]
        or cached.get("valid_from") != details["valid_from"].isoformat()
        or cached.get("valid_to") != details["valid_to"].isoformat()
        or cached.get("branch_id") != branch.id
        or cached.get("brand_id") != selected_brand.id
    ):
        raise ValidationError({"detail": "The selected Year and Season do not match the preview. Run preview again."})
    approved = [str(key) for key in group_keys if str(key)]
    groups = {group["import_key"]: group for group in cached["importable"]}
    unknown = [key for key in approved if key not in groups]
    if unknown:
        raise ValidationError({"group_keys": "Approved price list group was not part of the preview. Run preview again."})
    to_create = [groups[key] for key in approved]
    brands = _brand_map()
    stale = []
    for group in to_create:
        brand = brands.get(group["brand_code"].lower())
        if not brand or brand.id != group["brand_id"]:
            stale.append(f"Brand code {group['brand_code']} changed since preview")
            continue
        seen: set[tuple[str, str]] = set()
        for item in group["entries"]:
            record = {"brand_code": group["brand_code"], "product_code": item["product_code"], "season_label": group["season_label"], "valid_from": group["valid_from"], "valid_to": group["valid_to"], "price_basis": item["price_basis"], "taka_length_meters": item["taka_length_meters"], "price_value": item["price_value"]}
            current_brand, product, reason = _row_reason(record, brands, seen)
            if reason or not product or product.id != item["product_id"]:
                stale.append(f"Product code {item['product_code']} in {group['season_label']} changed since preview")
            seen.add((_group_key(record), item["product_code"].lower()))
    if stale:
        raise ValidationError({"detail": "Import was rejected because data changed since preview. " + " ".join(stale)})

    created = []
    with transaction.atomic():
        for group in to_create:
            price_list = PriceList.objects.create(brand_id=group["brand_id"], season_label=group["season_label"], valid_from=group["valid_from"], valid_to=group["valid_to"], status=PriceList.Status.DRAFT, created_by=user)
            for item in group["entries"]:
                PriceListEntry.objects.create(price_list=price_list, product_id=item["product_id"], taka_length_meters=Decimal(item["taka_length_meters"]), price_basis=item["price_basis"], price_value=Decimal(item["price_value"]), taka_quantity=item["taka_quantity"], status=PriceListEntry.Status.ACTIVE)
            created.append({"id": price_list.id, "season_label": price_list.season_label, "entry_count": len(group["entries"])})
    drop_preview(token)
    return {"created": created, "created_count": len(created), "entry_count": sum(item["entry_count"] for item in created), "excluded_count": len(cached["excluded"])}


def drop_preview(token: str) -> None:
    if token:
        cache.delete(f"{PREVIEW_CACHE_PREFIX}{token}")
