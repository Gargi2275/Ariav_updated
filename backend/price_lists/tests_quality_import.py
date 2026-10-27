from decimal import Decimal
from io import BytesIO

from django.core.cache import cache
from django.test import TestCase
from openpyxl import Workbook
from rest_framework.test import APIClient

from accounts.models import AuthUser
from brands.models import Brand
from categories.models import Category
from entities.models import Entity

from .models import PriceList, PriceListImportBatch, PriceListQualityEntry, PriceListQualityShade


class OcmQualityImportTests(TestCase):
    def setUp(self):
        cache.clear()
        self.user = AuthUser.objects.create_user(username="ocm-quality-admin", password="password", role=AuthUser.Role.ADMIN)
        self.brand = Brand.objects.create(brand_code="OCM", brand_name="OSM India Limited", order_method="Digital PO")
        self.branch = Entity.objects.create(short_code="OCM-BR", entity_name="OCM Branch", entity_type="Branch", status="Active")
        self.branch.brands.add(self.brand)
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def workbook(self, extra=False):
        book = Workbook()
        sheet = book.active
        sheet.title = "Price List"
        sheet.append(["Regular Price List-Winter'2026", None, None, None, None, None, None])
        sheet.append(["w.e.f 03.06.2026", None, None, None, "INR Per Mtr.*", None, None])
        sheet.append(["Serial No.", "Quality No.", "Shades", "Delivery", "Ex-Mill", "W/S", "Retail"])
        sheet.append([None, "Woolrite (AI)", None, None, None, None, None])
        sheet.append([1, "AI-1-2-209-0100100", "1,2,3,4,5,6,7,9,10,11,12,13,14,15,16,", "June", "1,799", "2,375", "-"])
        sheet.append([2, "AI-1-2-209-0100101", "1,2,3,4,5,6,7,8,", "June", "1,799", "2,375", "-"])
        sheet.append([1, "Ambassador White (AMB)", None, None, None, None, None])
        sheet.append([1, "AMB-1-1-999-0055101", "0A1", "October", "999", "1,319", "1,888"])
        sheet.append([2, "AMB-1-1-999-0055102", "0A1", "October", "999", "1,319", "1,888"])
        sheet.append(["Page 2 of 19", None, None, None, None, None, None])
        sheet.append(["Serial No.", "Quality No.", "Shades", "Delivery", "Ex-Mill", "W/S", "Retail"])
        sheet.append([None, "Aristo (ARO)", None, None, None, None, None])
        shades = ",".join(str(index) for index in range(1, 69)) + ","
        split_at = shades.find(",", 40) + 1
        sheet.append([1, "ARO-1-1-105-0041000", shades[:split_at], "June", "1,299", "1,728", "-"])
        sheet.append([None, None, shades[split_at:], None, None, None, None])
        if extra:
            sheet.cell(row=3, column=8, value="Unknown Column")
            for row in range(4, 15):
                sheet.cell(row=row, column=8, value="extra")
        output = BytesIO()
        book.save(output)
        return output.getvalue()

    def preview(self, content=None):
        return self.client.post(
            "/api/price-lists/import/preview/",
            {
                "file": BytesIO(content or self.workbook()),
                "format_id": "ocm_regular_pricelist",
                "branch_id": self.branch.id,
                "brand_id": self.brand.id,
            },
            format="multipart",
        )

    def test_sample_layout_is_fully_previewed(self):
        response = self.preview()
        self.assertEqual(response.status_code, 200, response.content)
        body = response.json()
        self.assertEqual((body["season"], body["year"]), ("Winter", 2026))
        self.assertEqual((body["valid_from"], body["valid_to"]), ("2026-06-03", "2026-12-31"))
        self.assertEqual(body["collections"][0]["code"], "AI")
        self.assertEqual(body["skipped_lines"]["page"], 1)
        self.assertEqual(body["skipped_lines"]["header"], 2)
        aro = next(row for row in body["entries"] if row["quality_no"] == "ARO-1-1-105-0041000")
        self.assertEqual(len(aro["shade_tokens"]), 68)
        self.assertNotIn("", aro["shade_tokens"])
        ai = next(row for row in body["entries"] if row["quality_no"] == "AI-1-2-209-0100100")
        self.assertIsNone(ai["retail"])
        ambassador = next(row for row in body["entries"] if row["quality_no"] == "AMB-1-1-999-0055101")
        self.assertEqual(ambassador["wholesale"], "1319")
        self.assertTrue(any(flag["code"] == "missing_retail" for flag in ai["flags"]))
        self.assertIn("0A1", ambassador["shade_tokens"])

    def test_errors_block_commit_and_warnings_require_confirmation(self):
        response = self.preview(self.workbook(extra=True))
        self.assertEqual(response.status_code, 200)
        body = response.json()
        # The sample has warnings for missing Retail but no errors.
        commit = self.client.post(
            "/api/price-lists/import/commit/",
            {
                "format_id": "ocm_regular_pricelist",
                "preview_token": body["preview_token"],
                "group_keys": [],
                "branch_id": self.branch.id,
                "brand_id": self.brand.id,
            },
            format="json",
        )
        self.assertEqual(commit.status_code, 400)
        self.assertIn("Confirm", commit.json()["detail"])

    def test_commit_persists_batch_entries_and_shades(self):
        body = self.preview().json()
        response = self.client.post(
            "/api/price-lists/import/commit/",
            {
                "format_id": "ocm_regular_pricelist",
                "preview_token": body["preview_token"],
                "group_keys": [],
                "branch_id": self.branch.id,
                "brand_id": self.brand.id,
                "confirmed_warnings": True,
            },
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(Category.objects.count(), 0)
        self.assertEqual(PriceList.objects.count(), 1)
        self.assertEqual(PriceListImportBatch.objects.count(), 1)
        self.assertEqual(PriceListQualityEntry.objects.count(), 5)
        self.assertEqual(PriceListQualityShade.objects.filter(entry__quality_no="ARO-1-1-105-0041000").count(), 68)
        self.assertEqual(PriceListQualityEntry.objects.get(quality_no="AMB-1-1-999-0055101").wholesale, Decimal("1319.00"))

    def test_same_file_is_rejected_after_commit(self):
        content = self.workbook()
        body = self.preview(content).json()
        commit = self.client.post(
            "/api/price-lists/import/commit/",
            {
                "format_id": "ocm_regular_pricelist",
                "preview_token": body["preview_token"],
                "group_keys": [],
                "branch_id": self.branch.id,
                "brand_id": self.brand.id,
                "confirmed_warnings": True,
            },
            format="json",
        )
        self.assertEqual(commit.status_code, 201, commit.content)
        response = self.preview(content)
        self.assertEqual(response.status_code, 400)
        self.assertIn("already imported", response.json()["file"])
