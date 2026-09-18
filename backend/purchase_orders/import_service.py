"""Purchase Order bulk import (Manual/POR) — preview is read-only; commit writes.

Customer codes are resolved only against the selected entity. Missing customers
are never auto-created. Rates stay null until staff enter them on the Draft PO.
"""

from __future__ import annotations

import hashlib
import secrets
from collections import OrderedDict
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.core.cache import cache
from django.db import transaction
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

from brands.models import Brand
from customers.models import CustomerEntity
from entities.models import Entity
from purchase_orders.models import PurchaseOrder, PurchaseOrderLine

PREVIEW_CACHE_PREFIX = "po_import_preview:"
PREVIEW_TTL_SECONDS = 30 * 60
RATE_WARNING = (
    "Rate is not present in the source file — all lines will import "
    "with rate left blank. POs will be created as Draft and MUST have "
    "rates entered manually before they can be submitted."
)
MAX_DESCRIPTION = 240
HEADER_ALIASES = {
    "customer code": "customer_code",
    "customer name": "customer_name",
    "order no": "order_no",
    "order number": "order_no",
    "order no.": "order_no",
    "code1": "code1",
    "code 1": "code1",
    "code5": "code5",
    "code 5": "code5",
    "code6": "code6",
    "code 6": "code6",
    "template": "template",
    "size": "size",
    "ord qty": "qty",
    "ord.qty": "qty",
    "order qty": "qty",
    "qty": "qty",
    "quantity": "qty",
    "agent code": "agent_code",
}


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).strip()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _norm_header(value) -> str:
    raw = _cell_str(value).lower().replace("_", " ").replace(".", " ")
    return " ".join(raw.split())


def compose_description(parts: dict) -> str:
    chunks = [
        _cell_str(parts.get("code1")),
        _cell_str(parts.get("code5")),
        _cell_str(parts.get("code6")),
        _cell_str(parts.get("template")),
        _cell_str(parts.get("size")),
    ]
    text = "-".join(chunks)
    if len(text) > MAX_DESCRIPTION:
        return text[:MAX_DESCRIPTION]
    return text


def _parse_qty(value) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, Decimal):
        qty = value
    elif isinstance(value, int):
        qty = Decimal(value)
    elif isinstance(value, float):
        qty = Decimal(str(value))
    else:
        raw = _cell_str(value).replace(",", "")
        if not raw:
            return None
        try:
            qty = Decimal(raw)
        except InvalidOperation:
            return None
    if qty <= 0:
        return None
    return qty.quantize(Decimal("0.01"))


def parse_workbook(file_bytes: bytes) -> tuple[list[dict], list[str]]:
    """Return (order groups in file order, row-level warning strings)."""
    warnings: list[str] = []
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError({"file": f"Could not read the spreadsheet: {exc}"}) from exc

    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            raise ValidationError({"file": "The spreadsheet is empty."}) from None

        mapping: dict[int, str] = {}
        for index, cell in enumerate(header_row or ()):
            key = HEADER_ALIASES.get(_norm_header(cell))
            if key:
                mapping[index] = key

        mapped_values = set(mapping.values())
        missing = []
        if "customer_code" not in mapped_values:
            missing.append("Customer Code")
        if "order_no" not in mapped_values:
            missing.append("Order No")
        if "qty" not in mapped_values:
            missing.append("Ord.Qty")
        if missing:
            raise ValidationError(
                {"file": f"Missing required column(s): {', '.join(missing)}."}
            )

        grouped: OrderedDict[str, dict] = OrderedDict()
        for excel_row, values in enumerate(rows, start=2):
            if not values or all(v is None or _cell_str(v) == "" for v in values):
                continue
            record = {field: None for field in (
                "customer_code", "customer_name", "order_no",
                "code1", "code5", "code6", "template", "size", "qty",
            )}
            for index, field in mapping.items():
                if index < len(values):
                    record[field] = values[index]

            order_no = _cell_str(record["order_no"])
            if not order_no:
                warnings.append(f"Row {excel_row}: skipped — Order No is blank.")
                continue

            customer_code = _cell_str(record["customer_code"])
            customer_name = _cell_str(record["customer_name"])
            qty = _parse_qty(record["qty"])
            if qty is None:
                warnings.append(
                    f"Row {excel_row} (Order No {order_no}): skipped line — Ord.Qty is missing or invalid."
                )
                if order_no not in grouped:
                    grouped[order_no] = {
                        "order_no": order_no,
                        "customer_codes": [],
                        "customer_name": customer_name,
                        "lines": [],
                    }
                continue

            description = compose_description(record)
            group = grouped.get(order_no)
            if group is None:
                grouped[order_no] = {
                    "order_no": order_no,
                    "customer_codes": [customer_code] if customer_code else [],
                    "customer_name": customer_name,
                    "lines": [{"description": description, "quantity": str(qty), "rate": None}],
                }
            else:
                if customer_code and customer_code not in group["customer_codes"]:
                    group["customer_codes"].append(customer_code)
                if customer_name and not group["customer_name"]:
                    group["customer_name"] = customer_name
                group["lines"].append(
                    {"description": description, "quantity": str(qty), "rate": None}
                )
    finally:
        workbook.close()

    return list(grouped.values()), warnings


def _customers_for_entity(entity: Entity) -> dict[str, object]:
    links = (
        CustomerEntity.objects.filter(entity=entity)
        .select_related("customer")
    )
    by_code: dict[str, object] = {}
    for link in links:
        code = (link.customer.customer_code or "").strip().upper()
        if code:
            by_code[code] = link.customer
    return by_code


def _existing_po_numbers() -> dict[str, str]:
    mapping: dict[str, str] = {}
    for number in PurchaseOrder.objects.values_list("po_number", flat=True):
        mapping[number.strip().lower()] = number
    return mapping


def unmatched_reason(customer_code: str, entity: Entity) -> str:
    code = customer_code or "(blank)"
    return f"Customer code {code} not found under {entity.entity_name}"


def build_preview(file_bytes: bytes, entity: Entity, brand: Brand) -> dict:
    if brand.order_method != Brand.OrderMethod.MANUAL_POR:
        raise ValidationError(
            {"brand_id": "Import is only available for Manual/POR brands."}
        )

    groups, parse_warnings = parse_workbook(file_bytes)
    customers = _customers_for_entity(entity)
    existing = _existing_po_numbers()

    importable: list[dict] = []
    excluded: list[dict] = []
    warnings = [RATE_WARNING, *parse_warnings]

    for group in groups:
        order_no = group["order_no"]
        codes = [c for c in group["customer_codes"] if c]
        line_count = len(group["lines"])
        customer_name = group["customer_name"]
        customer_code = codes[0] if codes else ""

        if line_count == 0:
            excluded.append(
                {
                    "order_no": order_no,
                    "customer_code": customer_code,
                    "customer_name": customer_name,
                    "reason": "No valid line items",
                    "line_count": 0,
                }
            )
            continue

        if len(codes) > 1:
            excluded.append(
                {
                    "order_no": order_no,
                    "customer_code": ", ".join(codes),
                    "customer_name": customer_name,
                    "reason": "Order No has mixed customer codes",
                    "line_count": line_count,
                }
            )
            warnings.append(
                f"Order No {order_no} has mixed customer codes ({', '.join(codes)}) and will not import."
            )
            continue

        collision = existing.get(order_no.lower())
        if collision:
            reason = f"Order No already exists as an existing po_number"
            excluded.append(
                {
                    "order_no": order_no,
                    "customer_code": customer_code,
                    "customer_name": customer_name,
                    "reason": reason,
                    "line_count": line_count,
                }
            )
            warnings.append(f"Order No {order_no} already exists as an existing po_number.")
            continue

        if not customer_code:
            excluded.append(
                {
                    "order_no": order_no,
                    "customer_code": "",
                    "customer_name": customer_name,
                    "reason": unmatched_reason("(blank)", entity),
                    "line_count": line_count,
                }
            )
            continue

        matched = customers.get(customer_code.upper())
        if matched is None:
            excluded.append(
                {
                    "order_no": order_no,
                    "customer_code": customer_code,
                    "customer_name": customer_name,
                    "reason": unmatched_reason(customer_code, entity),
                    "line_count": line_count,
                }
            )
            continue

        importable.append(
            {
                "order_no": order_no,
                "customer_code": customer_code,
                "customer_name": matched.customer_name,
                "file_customer_name": customer_name,
                "matched_customer_id": matched.id,
                "line_count": line_count,
                "lines": group["lines"],
            }
        )

    return {
        "entity_id": entity.id,
        "entity_name": entity.entity_name,
        "brand_id": brand.id,
        "file_sha256": hashlib.sha256(file_bytes).hexdigest(),
        "importable": importable,
        "excluded": excluded,
        "warnings": warnings,
    }


def store_preview(user_id: int, preview: dict) -> str:
    token = secrets.token_urlsafe(32)
    cache.set(
        f"{PREVIEW_CACHE_PREFIX}{token}",
        {"user_id": user_id, **preview},
        timeout=PREVIEW_TTL_SECONDS,
    )
    return token


def load_preview(token: str) -> dict | None:
    if not token:
        return None
    cached = cache.get(f"{PREVIEW_CACHE_PREFIX}{token}")
    return cached if isinstance(cached, dict) else None


def drop_preview(token: str) -> None:
    if token:
        cache.delete(f"{PREVIEW_CACHE_PREFIX}{token}")


def public_preview_payload(preview: dict, token: str) -> dict:
    importable = []
    for row in preview["importable"]:
        importable.append(
            {
                "order_no": row["order_no"],
                "customer_code": row["customer_code"],
                "customer_name": row["customer_name"],
                "matched_customer_id": row["matched_customer_id"],
                "line_count": row["line_count"],
                "lines": row["lines"],
            }
        )
    return {
        "preview_token": token,
        "entity_id": preview["entity_id"],
        "brand_id": preview["brand_id"],
        "importable": importable,
        "excluded": preview["excluded"],
        "warnings": preview["warnings"],
    }


def commit_import(*, user, token: str, entity_id: int, brand_id: int, order_nos: list[str]) -> dict:
    cached = load_preview(token)
    if not cached:
        raise ValidationError(
            {"detail": "Preview has expired or was not found. Run preview again before committing."}
        )
    if cached.get("user_id") != getattr(user, "id", None):
        raise ValidationError(
            {"detail": "Preview has expired or was not found. Run preview again before committing."}
        )
    if int(cached["entity_id"]) != int(entity_id):
        raise ValidationError({"entity_id": "entity_id does not match the preview."})
    if int(cached["brand_id"]) != int(brand_id):
        raise ValidationError({"brand_id": "brand_id does not match the preview."})

    approved = [_cell_str(value) for value in order_nos if _cell_str(value)]
    if not approved:
        raise ValidationError({"order_nos": "Select at least one purchase order to import."})

    importable_by_no = {row["order_no"]: row for row in cached["importable"]}
    excluded_by_no = {row["order_no"]: row for row in cached["excluded"]}
    unknown = [no for no in approved if no not in importable_by_no and no not in excluded_by_no]
    if unknown:
        raise ValidationError(
            {
                "order_nos": (
                    "Order No "
                    + ", ".join(unknown)
                    + " was not part of the preview. Run preview again."
                )
            }
        )

    to_create = [importable_by_no[no] for no in approved if no in importable_by_no]
    skipped_excluded = [excluded_by_no[no] for no in approved if no in excluded_by_no]

    entity = Entity.objects.filter(pk=entity_id).first()
    brand = Brand.objects.filter(pk=brand_id).first()
    if entity is None:
        raise ValidationError({"entity_id": "Entity was not found."})
    if brand is None:
        raise ValidationError({"brand_id": "Brand was not found."})
    if brand.order_method != Brand.OrderMethod.MANUAL_POR:
        raise ValidationError(
            {"brand_id": "Import is only available for Manual/POR brands."}
        )

    customers = _customers_for_entity(entity)
    existing = _existing_po_numbers()
    stale: list[str] = []

    for row in to_create:
        order_no = row["order_no"]
        if order_no.lower() in existing:
            stale.append(f"Order No {order_no} already exists as an existing po_number.")
            continue
        matched = customers.get((row["customer_code"] or "").upper())
        if matched is None:
            stale.append(
                f"Customer code {row['customer_code']} is no longer linked to {entity.entity_name}."
            )
            continue
        if matched.id != row["matched_customer_id"]:
            stale.append(
                f"Customer match for Order No {order_no} changed since preview. Run preview again."
            )

    if stale:
        raise ValidationError(
            {
                "detail": (
                    "Import was rejected because data changed since preview. "
                    + " ".join(stale)
                )
            }
        )

    created: list[dict] = []
    with transaction.atomic():
        for row in to_create:
            po = PurchaseOrder.objects.create(
                po_number=row["order_no"],
                order_type=PurchaseOrder.OrderType.MANUAL,
                entity=entity,
                customer_id=row["matched_customer_id"],
                brand=brand,
                po_date=date.today(),
                status=PurchaseOrder.Status.DRAFT,
                created_by=user if getattr(user, "is_authenticated", False) else None,
            )
            for line in row["lines"]:
                qty = Decimal(str(line["quantity"]))
                PurchaseOrderLine.objects.create(
                    purchase_order=po,
                    product=None,
                    product_description=(line.get("description") or "")[:MAX_DESCRIPTION],
                    quantity=qty,
                    rate=None,
                )
            created.append(
                {
                    "id": po.id,
                    "po_number": po.po_number,
                    "customer_id": po.customer_id,
                    "line_count": len(row["lines"]),
                }
            )

    drop_preview(token)
    remaining_excluded = list(cached["excluded"])
    return {
        "created": created,
        "created_count": len(created),
        "excluded": remaining_excluded,
        "excluded_count": len(remaining_excluded),
        "skipped_excluded_order_nos": [row["order_no"] for row in skipped_excluded],
    }
