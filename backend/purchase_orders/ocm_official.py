"""Official OCM order import, comparison, and approval.

The internal purchase order stays the planning record. An official order is
stored beside it and is linked only after a user approves the comparison.
Dispatch then matches Order Code to that official order number.
"""

from __future__ import annotations

import hashlib
import zipfile
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.core.files.base import ContentFile
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.exceptions import PermissionDenied, ValidationError

from accounts.access import can_delete_records, can_edit_records
from accounts.models import AuditLog
from brands.models import Brand
from customers.models import Customer
from purchase_orders.models import (
    OcmProductMapping,
    OfficialOcmAudit,
    OfficialOcmImport,
    OfficialOcmOrder,
    OfficialOcmOrderLine,
    PurchaseOrder,
)

MAX_FILE_BYTES = 5 * 1024 * 1024
CODE_FIELDS = ("code1", "code2", "code3", "code4", "code5", "code6", "template", "size")
HEADER_ALIASES = {
    "customercode": "customer_code",
    "customername": "customer_name",
    "orderno": "order_no",
    "ordernumber": "order_no",
    "season": "season",
    "code1": "code1",
    "code2": "code2",
    "code3": "code3",
    "code4": "code4",
    "code5": "code5",
    "code6": "code6",
    "template": "template",
    "size": "size",
    "ordqty": "ordered_qty",
    "orderqty": "ordered_qty",
    "dispqty": "dispatched_qty",
    "packedqty": "packed_qty",
    "alocqty": "allocated_qty",
    "allocqty": "allocated_qty",
    "pendingorderqty": "pending_qty",
    "dwliverystart": "delivery_start",
    "deliverystart": "delivery_start",
    "deliveryend": "delivery_end",
}
REQUIRED = ("customer_code", "order_no", "ordered_qty", "code1")


def ocm_brand() -> Brand:
    brand = Brand.objects.filter(brand_code__iexact="OCM").first()
    if brand is None:
        brand = Brand.objects.filter(brand_name__iexact="OSM India Limited").first()
    if brand is None:
        raise ValidationError({"brand": "OCM is not configured as a brand."})
    return brand


def _cell(value) -> str:
    """Cell text. Strings keep leading zeros and hyphens; only numeric cells are shortened."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, bool):
        return str(value).strip()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header_key(value) -> str:
    return "".join(ch for ch in _cell(value).casefold() if ch.isalnum())


def identity_key(parts: dict) -> str:
    return "\x1f".join(_cell(parts.get(name)) for name in CODE_FIELDS)


def _qty(value, *, required: bool) -> Decimal | None:
    text = _cell(value).replace(",", "")
    if text == "":
        if required:
            raise ValueError("blank")
        return None
    try:
        quantity = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("invalid") from exc
    if quantity < 0:
        raise ValueError("negative")
    return quantity.quantize(Decimal("0.01"))


def _parse_date(value):
    if value is None or _cell(value) == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        serial = int(value)
        if serial <= 0 or serial > 60000:
            raise ValueError("invalid date")
        return (datetime(1899, 12, 30) + timedelta(days=serial)).date()
    text = _cell(value)[:10]
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError("invalid date")


def _workbook_kind(filename: str) -> str:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        return "xlsx"
    if name.endswith(".xlsb"):
        return "xlsb"
    return ""


def _zip_names(file_bytes: bytes) -> set[str]:
    try:
        with zipfile.ZipFile(BytesIO(file_bytes)) as archive:
            return {item.replace("\\", "/").lower() for item in archive.namelist()}
    except zipfile.BadZipFile as exc:
        raise ValidationError({
            "file": "The file is not a readable Excel workbook. It may be corrupt or renamed.",
        }) from exc


def _sheet_has_headers(rows) -> bool:
    for header in rows[:20]:
        _mapping, seen = _header_mapping(header)
        if all(name in seen for name in REQUIRED):
            return True
    return False


def _xlsx_rows(file_bytes: bytes):
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError({"file": f"Could not read the spreadsheet: {exc}"}) from exc
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        raise ValidationError({"file": "The spreadsheet is empty."})
    return rows


def _xlsb_rows(file_bytes: bytes):
    try:
        from pyxlsb import open_workbook
    except ImportError as exc:
        raise ValidationError({
            "file": "xlsb support is not installed. Install pyxlsb and import the file again.",
        }) from exc
    try:
        workbook = open_workbook(BytesIO(file_bytes))
    except Exception as exc:
        raise ValidationError({"file": f"Could not read the spreadsheet: {exc}"}) from exc
    try:
        if not workbook.sheets:
            raise ValidationError({"file": "The spreadsheet has no worksheets."})
        fallback = None
        for index in range(1, len(workbook.sheets) + 1):
            sheet = workbook.get_sheet(index)
            try:
                rows = [tuple(cell.v for cell in row) for row in sheet.rows()]
            finally:
                sheet.close()
            if fallback is None:
                fallback = rows
            if _sheet_has_headers(rows):
                return rows
        if not fallback:
            raise ValidationError({"file": "The spreadsheet is empty."})
        return fallback
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError({"file": f"Could not read the spreadsheet: {exc}"}) from exc
    finally:
        workbook.close()


def _workbook_rows(file_bytes: bytes, filename: str):
    kind = _workbook_kind(filename)
    if kind not in {"xlsx", "xlsb"}:
        raise ValidationError({"file": "Upload an .xlsx or .xlsb workbook."})
    if len(file_bytes) > MAX_FILE_BYTES:
        raise ValidationError({"file": "The file is larger than 5 MB."})
    if file_bytes[:8].startswith(b"\xd0\xcf\x11\xe0"):
        raise ValidationError({
            "file": (
                "This file is a legacy Excel workbook (.xls), not an .xlsx or .xlsb file. "
                "Save the Order Status workbook as .xlsx or .xlsb."
            ),
        })
    if not file_bytes.startswith(b"PK"):
        raise ValidationError({
            "file": "The file content is not an Excel workbook. A renamed or damaged file cannot be imported.",
        })
    names = _zip_names(file_bytes)
    has_xlsx = "xl/workbook.xml" in names
    has_xlsb = "xl/workbook.bin" in names
    if kind == "xlsx" and not has_xlsx:
        if has_xlsb:
            raise ValidationError({
                "file": "This file is an .xlsb workbook with an .xlsx name. Rename it to .xlsb or save it as .xlsx.",
            })
        raise ValidationError({"file": "The file is not a valid .xlsx workbook."})
    if kind == "xlsb" and not has_xlsb:
        if has_xlsx:
            raise ValidationError({
                "file": "This file is an .xlsx workbook with an .xlsb name. Rename it to .xlsx or save it as .xlsb.",
            })
        raise ValidationError({"file": "The file is not a valid .xlsb workbook."})
    if kind == "xlsx":
        return _xlsx_rows(file_bytes)
    return _xlsb_rows(file_bytes)


def _header_mapping(header) -> tuple[dict[int, str], set[str]]:
    mapping: dict[int, str] = {}
    seen: set[str] = set()
    for index, cell in enumerate(header or ()):
        field = HEADER_ALIASES.get(_header_key(cell))
        if field and field not in seen:
            mapping[index] = field
            seen.add(field)
    return mapping, seen


def parse_official_workbook(file_bytes: bytes, filename: str) -> tuple[list[dict], list[dict]]:
    rows = _workbook_rows(file_bytes, filename)
    header_offset = None
    mapping: dict[int, str] = {}
    seen: set[str] = set()
    for offset, header in enumerate(rows[:20]):
        mapping, seen = _header_mapping(header)
        if all(name in seen for name in REQUIRED):
            header_offset = offset
            break
    if header_offset is None:
        labels = {
            "customer_code": "Customer Code",
            "order_no": "Order No",
            "ordered_qty": "Ord.Qty",
            "code1": "Code1",
        }
        missing = [labels[name] for name in REQUIRED if name not in seen]
        raise ValidationError({
            "file": "Missing required column(s): " + ", ".join(missing) + ".",
        })

    orders: dict[str, dict] = {}
    errors: list[dict] = []
    for row_number, values in enumerate(rows[header_offset + 1:], start=header_offset + 2):
        if not values or all(_cell(value) == "" for value in values):
            continue
        record = {field: "" for field in HEADER_ALIASES.values()}
        for index, field in mapping.items():
            if index < len(values):
                record[field] = values[index]
        order_no = _cell(record["order_no"])
        if not order_no:
            errors.append({"row": row_number, "reason": "Order No is blank."})
            continue
        try:
            ordered = _qty(record["ordered_qty"], required=True)
            dispatched = _qty(record.get("dispatched_qty"), required=False)
            packed = _qty(record.get("packed_qty"), required=False)
            allocated = _qty(record.get("allocated_qty"), required=False)
            pending = _qty(record.get("pending_qty"), required=False)
            delivery_start = _parse_date(record.get("delivery_start"))
            delivery_end = _parse_date(record.get("delivery_end"))
        except ValueError as exc:
            errors.append({"row": row_number, "reason": f"Could not read quantities or dates ({exc})."})
            continue
        bucket = orders.setdefault(order_no, {
            "order_number": order_no,
            "customer_code": _cell(record["customer_code"]),
            "customer_name": _cell(record["customer_name"]),
            "season": _cell(record.get("season")),
            "lines": [],
        })
        if _cell(record["customer_code"]) and bucket["customer_code"] and _cell(record["customer_code"]).casefold() != bucket["customer_code"].casefold():
            errors.append({"row": row_number, "reason": "Customer Code changes inside the same Order No."})
            continue
        parts = {name: _cell(record.get(name)) for name in CODE_FIELDS}
        bucket["lines"].append({
            "row_number": row_number,
            **parts,
            "identity_key": identity_key(parts),
            "ordered_qty": ordered,
            "dispatched_qty": dispatched,
            "packed_qty": packed,
            "allocated_qty": allocated,
            "pending_qty": pending,
            "delivery_start": delivery_start,
            "delivery_end": delivery_end,
            "raw": {key: _cell(value) if not isinstance(value, (date, datetime, Decimal)) else str(value) for key, value in record.items()},
        })
    return list(orders.values()), errors


def _customer_for(code: str):
    if not code:
        return None, "Customer Code is blank."
    matches = list(Customer.objects.filter(customer_code__iexact=code.strip()))
    if len(matches) == 1:
        return matches[0], ""
    if not matches:
        return None, f"Customer Code {code} was not found."
    return None, f"Customer Code {code} matches more than one party."


def _refresh_status(order: OfficialOcmOrder) -> None:
    if order.status == OfficialOcmOrder.Status.APPROVED:
        return
    report = compare_order(order)
    order.status = (
        OfficialOcmOrder.Status.NEEDS_REVIEW
        if report["blockers"]
        else OfficialOcmOrder.Status.READY
        if order.candidate_purchase_order_id or len(report["candidates"]) == 1
        else OfficialOcmOrder.Status.AWAITING_REVIEW
    )
    order.save(update_fields=["status"])


@transaction.atomic
def import_official_file(file_bytes: bytes, filename: str, user) -> dict:
    brand = ocm_brand()
    if brand.brand_code.casefold() != "ocm" and brand.brand_name.casefold() != "osm india limited":
        raise ValidationError({"brand": "Official order import is only available for OCM."})
    checksum = hashlib.sha256(file_bytes).hexdigest()
    existing = OfficialOcmImport.objects.filter(checksum=checksum).first()
    if existing:
        return {
            "duplicate_file": True,
            "import_id": existing.id,
            "orders": [_order_summary(order) for order in existing.orders.all()],
            "row_errors": [],
        }
    parsed, row_errors = parse_official_workbook(file_bytes, filename)
    if not parsed and row_errors:
        raise ValidationError({"rows": row_errors})
    batch = OfficialOcmImport(
        brand=brand,
        filename=(filename or "official.xlsx")[:240],
        checksum=checksum,
        imported_by=user if getattr(user, "is_authenticated", False) else None,
    )
    batch.source_file.save(batch.filename, ContentFile(file_bytes), save=False)
    batch.save()
    created = []
    for item in parsed:
        if OfficialOcmOrder.objects.filter(brand=brand, order_number=item["order_number"]).exists():
            row_errors.append({
                "row": item["lines"][0]["row_number"] if item["lines"] else None,
                "reason": f"Official order {item['order_number']} is already imported.",
            })
            continue
        customer, customer_error = _customer_for(item["customer_code"])
        order = OfficialOcmOrder.objects.create(
            import_batch=batch,
            brand=brand,
            customer=customer,
            order_number=item["order_number"],
            season=item["season"],
            customer_code=item["customer_code"],
            customer_name=item["customer_name"],
            status=OfficialOcmOrder.Status.NEEDS_REVIEW if customer_error else OfficialOcmOrder.Status.AWAITING_REVIEW,
        )
        seen: dict[str, int] = {}
        for line in item["lines"]:
            duplicate = line["identity_key"] in seen
            if duplicate:
                OfficialOcmOrderLine.objects.filter(pk=seen[line["identity_key"]]).update(duplicate=True)
            stored = OfficialOcmOrderLine.objects.create(order=order, duplicate=duplicate, **line)
            seen[line["identity_key"]] = stored.id
        if customer_error or any(line.duplicate for line in order.lines.all()):
            order.status = OfficialOcmOrder.Status.NEEDS_REVIEW
            order.save(update_fields=["status"])
        else:
            _refresh_status(order)
        created.append(_order_summary(order))
    return {"duplicate_file": False, "import_id": batch.id, "orders": created, "row_errors": row_errors}


def _order_summary(order: OfficialOcmOrder) -> dict:
    return {
        "id": order.id,
        "order_number": order.order_number,
        "status": order.status,
        "customer_code": order.customer_code,
        "customer_name": order.customer.customer_name if order.customer_id else order.customer_name,
        "filename": order.import_batch.filename,
        "imported_at": order.import_batch.imported_at.isoformat() if order.import_batch_id else "",
        "purchase_order_id": order.purchase_order_id,
        "purchase_order_number": order.purchase_order.po_number if order.purchase_order_id else "",
    }


def _line_identity_label(line) -> str:
    return " / ".join(part for part in (_cell(getattr(line, name)) for name in CODE_FIELDS) if part)


def _full_identity(official_line) -> str:
    return "-".join(_cell(getattr(official_line, name)) for name in ("code1", "code5", "code6", "template", "size"))


def _short_identity(code1: str, code5: str, size: str, shade: str):
    if not code1 or not code5:
        return None
    return (code1.casefold(), code5.casefold(), size.casefold(), shade.casefold())


def _official_short_identity(official_line):
    return _short_identity(
        _cell(official_line.code1),
        _cell(official_line.code5),
        _cell(official_line.size),
        _cell(official_line.code2),
    )


def _internal_short_identity(po_line):
    parts = _cell(po_line.product_description).split("-")
    if len(parts) != 2:
        return None
    return _short_identity(parts[0], parts[1], _cell(po_line.size), _cell(po_line.shade))


def _internal_match(po_line, official_line, mapping: OcmProductMapping | None) -> bool:
    if mapping and mapping.product_id and po_line.product_id == mapping.product_id:
        if mapping.shade and _cell(po_line.shade).casefold() != mapping.shade.strip().casefold():
            return False
        return True
    if mapping and mapping.quality_label:
        if _cell(po_line.product_description).casefold() != mapping.quality_label.strip().casefold():
            return False
        if mapping.shade and _cell(po_line.shade).casefold() != mapping.shade.strip().casefold():
            return False
        return True
    composed = _full_identity(official_line)
    return bool(composed.strip("-")) and _cell(po_line.product_description) == composed


def compare_order(order: OfficialOcmOrder) -> dict:
    brand = order.brand
    candidates = list(
        PurchaseOrder.objects.filter(brand=brand, customer=order.customer)
        .select_related("customer", "entity", "brand")
        .prefetch_related("lines")
    ) if order.customer_id else []
    chosen = order.purchase_order or order.candidate_purchase_order
    if chosen and chosen.brand_id == brand.id:
        pool = [chosen]
    else:
        exact = [
            po for po in candidates
            if _cell(po.brand_order_code).casefold() == order.order_number.casefold()
        ]
        open_rows = [po for po in candidates if not _cell(po.brand_order_code)]
        conflicting = [
            po for po in candidates
            if _cell(po.brand_order_code) and _cell(po.brand_order_code).casefold() != order.order_number.casefold()
        ]
        pool = exact or open_rows
        if conflicting and not exact:
            pool = []
    mappings = {
        item.identity_key: item
        for item in OcmProductMapping.objects.filter(brand=brand, approved=True)
    }
    blockers: list[str] = []
    if order.customer_id is None:
        blockers.append("The official customer code is not a single known party.")
    if len(pool) == 0:
        blockers.append("No internal OCM purchase order can be matched without conflicting with another official order number.")
    elif len(pool) > 1 and chosen is None:
        blockers.append("More than one internal purchase order matches. Select one.")
    po = pool[0] if len(pool) == 1 else None
    if po and order.customer_id and po.customer_id != order.customer_id:
        blockers.append("The internal purchase order belongs to a different party.")
        po = None
    comparisons = []
    used_internal: set[int] = set()
    if any(line.duplicate for line in order.lines.all()):
        blockers.append("The official file contains a duplicate product identity.")
    if po:
        internal_lines = list(po.lines.all())
        official_lines = list(order.lines.all())
        official_short: dict[tuple, int] = {}
        internal_short: dict[tuple, int] = {}
        for official in official_lines:
            key = _official_short_identity(official)
            if key:
                official_short[key] = official_short.get(key, 0) + 1
        for internal in internal_lines:
            key = _internal_short_identity(internal)
            if key:
                internal_short[key] = internal_short.get(key, 0) + 1
        for official in official_lines:
            mapping = mappings.get(official.identity_key)
            matches = [
                line for line in internal_lines
                if line.id not in used_internal and _internal_match(line, official, mapping)
            ]
            short_key = _official_short_identity(official)
            automatic = False
            if mapping is None and not matches and short_key and official_short.get(short_key) == 1 and internal_short.get(short_key) == 1:
                matches = [
                    line for line in internal_lines
                    if line.id not in used_internal and _internal_short_identity(line) == short_key
                ]
                automatic = bool(matches)
            if official.duplicate:
                status = "Duplicate Candidate"
                reason = "The same OCM product identity appears more than once in the official file."
                internal = None
            elif mapping is None and not matches and short_key and (official_short.get(short_key, 0) > 1 or internal_short.get(short_key, 0) > 1):
                status = "Needs Review"
                reason = "Several lines share this code, shade, and size, so none was selected automatically."
                internal = None
                blockers.append(f"Row {official.row_number} has more than one possible line. Choose the mapping.")
            elif mapping is None and not matches:
                status = "Unmatched"
                reason = "No approved mapping exists for this OCM product identity."
                internal = None
                blockers.append(f"Row {official.row_number} has no approved product mapping.")
            elif len(matches) > 1:
                status = "Needs Review"
                reason = "More than one internal line matches this product identity."
                internal = None
                blockers.append(f"Row {official.row_number} matches more than one internal line.")
            elif not matches:
                status = "Extra Official Line"
                reason = "The official order has a mapped product that is not on the internal purchase order."
                internal = None
                blockers.append(f"Row {official.row_number} is an extra official line.")
            else:
                internal = matches[0]
                used_internal.add(internal.id)
                status, reason = _difference(internal, official)
                if automatic:
                    reason = "Matched automatically because this code, shade, and size appear once. " + reason
            comparisons.append(_comparison_row(official, internal, status, reason))
        for internal in internal_lines:
            if internal.id in used_internal:
                continue
            comparisons.append({
                "official_line_id": None,
                "row_number": None,
                "internal_line_id": internal.id,
                "internal_product": internal.product.product_name if internal.product_id else internal.product_description,
                "internal_code": internal.product.product_code if internal.product_id else "",
                "ocm_identity": "",
                "template": "",
                "size": internal.size,
                "internal_qty": str(internal.quantity),
                "official_qty": "",
                "difference": "",
                "internal_rate": str(internal.rate) if internal.rate is not None else "",
                "official_rate": "",
                "internal_delivery": _window(internal.delivery_from, internal.delivery_to),
                "official_delivery": "",
                "delivery_start": "",
                "delivery_end": "",
                "code1": "",
                "code2": "",
                "code3": "",
                "code4": "",
                "code5": "",
                "code6": "",
                "status": "Missing in Official Order",
                "reason": "This internal line has no matching official line.",
            })
            blockers.append(f"Internal line {internal.id} is missing from the official order.")
    else:
        for official in order.lines.all():
            mapping = mappings.get(official.identity_key)
            status = "Duplicate Candidate" if official.duplicate else "Unmatched" if mapping is None else "Needs Review"
            reason = "Select one internal purchase order before the lines can be compared."
            comparisons.append(_comparison_row(official, None, status, reason))
    differences = [row for row in comparisons if row["status"] not in {"Exact Match"}]
    return {
        "order": _order_summary(order),
        "internal_po_number": po.po_number if po else "",
        "internal_po_id": po.id if po else None,
        "brand": brand.brand_code,
        "candidates": [{"id": item.id, "po_number": item.po_number} for item in pool],
        "lines": comparisons,
        "blockers": list(dict.fromkeys(blockers)),
        "differences": differences,
        "audit": [
            {
                "action": event.action,
                "reason": event.reason,
                "before": event.before,
                "after": event.after,
                "at": event.created_at.isoformat(),
                "user_id": event.created_by_id,
            }
            for event in order.audit_events.all()
        ],
    }


def _difference(internal, official) -> tuple[str, str]:
    reasons = []
    status = "Exact Match"
    if internal.quantity != official.ordered_qty:
        status = "Quantity Difference"
        reasons.append(
            f"Internal quantity {internal.quantity} and official ordered quantity {official.ordered_qty} differ."
        )
    if official.delivery_start and internal.delivery_from and official.delivery_start != internal.delivery_from:
        status = "Delivery Date Difference" if status == "Exact Match" else status
        reasons.append("Delivery start differs from the internal delivery from date.")
    if official.delivery_end and internal.delivery_to and official.delivery_end != internal.delivery_to:
        status = "Delivery Date Difference" if status == "Exact Match" else status
        reasons.append("Delivery end differs from the internal delivery to date.")
    return status, " ".join(reasons) or "Product identity, quantity, and delivery window match."


def _window(start, end) -> str:
    if not start and not end:
        return ""
    return f"{start.isoformat() if start else ''} – {end.isoformat() if end else ''}"


def _comparison_row(official, internal, status: str, reason: str) -> dict:
    return {
        "official_line_id": official.id,
        "row_number": official.row_number,
        "internal_line_id": internal.id if internal else None,
        "internal_product": (internal.product.product_name if internal and internal.product_id else internal.product_description if internal else ""),
        "internal_code": internal.product.product_code if internal and internal.product_id else "",
        "ocm_identity": _line_identity_label(official),
        "template": official.template,
        "size": official.size,
        "internal_qty": str(internal.quantity) if internal else "",
        "official_qty": str(official.ordered_qty),
        "difference": str((internal.quantity - official.ordered_qty) if internal else ""),
        "internal_rate": str(internal.rate) if internal and internal.rate is not None else "",
        "official_rate": "",
        "internal_delivery": _window(internal.delivery_from, internal.delivery_to) if internal else "",
        "official_delivery": _window(official.delivery_start, official.delivery_end),
        "delivery_start": official.delivery_start.isoformat() if official.delivery_start else "",
        "delivery_end": official.delivery_end.isoformat() if official.delivery_end else "",
        "code1": official.code1,
        "code2": official.code2,
        "code3": official.code3,
        "code4": official.code4,
        "code5": official.code5,
        "code6": official.code6,
        "status": status,
        "reason": reason,
    }


def select_internal_po(order: OfficialOcmOrder, purchase_order_id: int, user) -> dict:
    _require_editor(user)
    po = PurchaseOrder.objects.filter(pk=purchase_order_id, brand=order.brand).select_related("customer").first()
    if po is None:
        raise ValidationError({"purchase_order_id": "Choose an OCM purchase order."})
    if order.customer_id and po.customer_id != order.customer_id:
        raise ValidationError({"purchase_order_id": "That purchase order belongs to a different party."})
    if _cell(po.brand_order_code) and _cell(po.brand_order_code).casefold() != order.order_number.casefold():
        raise ValidationError({"purchase_order_id": "That purchase order is already tied to a different official order number."})
    if OfficialOcmOrder.objects.exclude(pk=order.pk).filter(purchase_order=po).exists():
        raise ValidationError({"purchase_order_id": "That purchase order is already linked to an official OCM order."})
    before = order.candidate_purchase_order_id
    order.candidate_purchase_order = po
    order.save(update_fields=["candidate_purchase_order"])
    _refresh_status(order)
    OfficialOcmAudit.objects.create(
        order=order,
        action="select_po",
        created_by=user if getattr(user, "is_authenticated", False) else None,
        before={"candidate_purchase_order_id": before},
        after={"candidate_purchase_order_id": po.id, "po_number": po.po_number},
    )
    return compare_order(order)


def save_mapping(order: OfficialOcmOrder, payload: dict, user) -> dict:
    _require_editor(user)
    line = order.lines.filter(pk=payload.get("line_id")).first()
    if line is None:
        raise ValidationError({"line_id": "Choose an official line."})
    reason = _cell(payload.get("reason"))
    if not reason:
        raise ValidationError({"reason": "A reason is required to approve a product mapping."})
    product_id = payload.get("product_id")
    quality = _cell(payload.get("quality_label"))
    if not product_id and not quality:
        raise ValidationError({"quality_label": "Choose a product or enter the internal quality label."})
    mapping, _created = OcmProductMapping.objects.update_or_create(
        brand=order.brand,
        identity_key=line.identity_key,
        defaults={
            "code1": line.code1,
            "code2": line.code2,
            "code3": line.code3,
            "code4": line.code4,
            "code5": line.code5,
            "code6": line.code6,
            "template": line.template,
            "size": line.size,
            "product_id": product_id or None,
            "quality_label": quality,
            "shade": _cell(payload.get("shade")),
            "approved": True,
            "reason": reason,
            "created_by": user if getattr(user, "is_authenticated", False) else None,
        },
    )
    OfficialOcmAudit.objects.create(
        order=order,
        action="mapping",
        reason=reason,
        created_by=mapping.created_by,
        before={},
        after={"identity_key": line.identity_key, "product_id": mapping.product_id, "quality_label": mapping.quality_label},
    )
    _refresh_status(order)
    return compare_order(order)


@transaction.atomic
def approve_order(order_id: int, reason: str, user) -> dict:
    _require_editor(user)
    order = OfficialOcmOrder.objects.select_for_update().select_related("brand", "customer", "purchase_order").get(pk=order_id)
    if order.status == OfficialOcmOrder.Status.APPROVED:
        raise ValidationError({"detail": "This official order is already approved."})
    report = compare_order(order)
    if report["blockers"]:
        raise ValidationError({"detail": "Approval is blocked until every critical conflict is resolved.", "blockers": report["blockers"]})
    if report["differences"] and not reason.strip():
        raise ValidationError({"reason": "A reason is required when the official order differs from the internal purchase order."})
    po = PurchaseOrder.objects.select_for_update().get(pk=report["internal_po_id"])
    original_number = po.po_number
    original_code = po.brand_order_code
    if po.brand_id != order.brand_id:
        raise ValidationError({"detail": "The internal purchase order is not an OCM order."})
    if original_code and original_code.casefold() != order.order_number.casefold():
        raise ValidationError({"detail": "The internal purchase order already has a different official order number."})
    po.brand_order_code = order.order_number
    po.save(update_fields=["brand_order_code", "updated_at"])
    order.purchase_order = po
    order.status = OfficialOcmOrder.Status.APPROVED
    order.approved_by = user if getattr(user, "is_authenticated", False) else None
    order.approved_at = timezone.now()
    order.save(update_fields=["purchase_order", "status", "approved_by", "approved_at"])
    OfficialOcmAudit.objects.create(
        order=order,
        action="approve",
        reason=reason.strip(),
        created_by=order.approved_by,
        before={"po_number": original_number, "brand_order_code": original_code or ""},
        after={"po_number": po.po_number, "brand_order_code": po.brand_order_code},
    )
    AuditLog.objects.create(
        user=getattr(order.approved_by, "display_name", "") or getattr(order.approved_by, "username", ""),
        role=getattr(order.approved_by, "role", ""),
        action="Official OCM Order Approved",
        module="Purchase Orders",
        details=f"Official OCM order {order.order_number} linked to {po.po_number}.",
    )
    _mark_brand_accepted(po, order, order.approved_by)
    po.refresh_from_db()
    if po.po_number != original_number:
        raise ValidationError({"detail": "Approval must not change the internal purchase order number."})
    return compare_order(order)


def _mark_brand_accepted(po, order, user) -> None:
    """The approved official order is the brand's acceptance, so move the PO forward."""
    from .transitions import apply_status_change, validate_transition

    reason = f"Official OCM order {order.order_number} approved."
    steps = {
        PurchaseOrder.Status.SUBMITTED: [PurchaseOrder.Status.SENT_TO_BRAND, PurchaseOrder.Status.BRAND_ACCEPTED],
        PurchaseOrder.Status.SENT_TO_BRAND: [PurchaseOrder.Status.BRAND_ACCEPTED],
        PurchaseOrder.Status.ON_HOLD: [PurchaseOrder.Status.BRAND_ACCEPTED],
    }.get(po.status, [])
    for target in steps:
        if validate_transition(po.status, target):
            return
        apply_status_change(po, target, user=user, reason=reason)


def _drop_empty_import_batches(batch_ids) -> None:
    for batch in OfficialOcmImport.objects.filter(pk__in=batch_ids):
        if batch.orders.exists():
            continue
        if batch.source_file:
            batch.source_file.delete(save=False)
        batch.delete()


def _unapproved_orders():
    return OfficialOcmOrder.objects.filter(purchase_order__isnull=True).exclude(
        status=OfficialOcmOrder.Status.APPROVED,
    )


@transaction.atomic
def delete_official_orders(order_ids, user, *, all_unapproved: bool = False) -> dict:
    """Remove official OCM orders. Internal purchase orders and dispatch rows stay."""
    if not can_delete_records(user):
        raise PermissionDenied("You do not have permission to delete official OCM orders.")
    blocked = []
    missing: list[int] = []
    if all_unapproved:
        eligible = _unapproved_orders()
        deleted_count = eligible.count()
        batch_ids = list(eligible.values_list("import_batch_id", flat=True).distinct())
        eligible.delete()
        deleted = []
    else:
        if not isinstance(order_ids, (list, tuple)) or not order_ids:
            raise ValidationError({"ids": "Choose at least one official order."})
        if len(order_ids) > 500:
            raise ValidationError({"ids": "Delete at most 500 official orders at a time."})
        parsed: list[int] = []
        for raw in order_ids:
            try:
                parsed.append(int(raw))
            except (TypeError, ValueError) as exc:
                raise ValidationError({"ids": "Each id must be an official order id."}) from exc
        found = set(OfficialOcmOrder.objects.filter(pk__in=parsed).values_list("pk", flat=True))
        missing = [item for item in parsed if item not in found]
        blocked = [
            {
                "id": order.id,
                "order_number": order.order_number,
                "reason": "This official order is approved and stays linked to its internal purchase order.",
            }
            for order in OfficialOcmOrder.objects.filter(pk__in=parsed).exclude(pk__in=_unapproved_orders().values("pk"))
        ]
        eligible = _unapproved_orders().filter(pk__in=parsed)
        deleted = [{"id": pk, "order_number": number} for pk, number in eligible.values_list("id", "order_number")]
        deleted_count = len(deleted)
        batch_ids = list(eligible.values_list("import_batch_id", flat=True).distinct())
        eligible.delete()
    _drop_empty_import_batches(batch_ids)
    remaining = _unapproved_orders().count()
    return {
        "deleted": deleted,
        "deleted_count": deleted_count,
        "blocked": blocked,
        "missing": missing,
        "remaining": remaining,
    }


def _require_editor(user):
    if not can_edit_records(user):
        raise PermissionDenied("You do not have permission to change official OCM orders.")


def dispatch_target_for_order_code(order_code: str, visible_ids: set[int] | None = None):
    """When an official OCM order exists, dispatch may use only its approved purchase order.

    Returns (purchase_order, error). A missing official order returns (None, None)
    so the existing dispatch matcher can keep its current behaviour.
    """
    order = (
        OfficialOcmOrder.objects.select_related("purchase_order", "brand")
        .filter(order_number__iexact=_cell(order_code), brand__brand_code__iexact="OCM")
        .first()
    )
    if order is None:
        return None, None
    if order.status != OfficialOcmOrder.Status.APPROVED or order.purchase_order_id is None:
        return None, (
            f"Official OCM order {order.order_number} is not approved, so this dispatch was not attached."
        )
    if visible_ids is not None and order.purchase_order_id not in visible_ids:
        return None, f"No matching Purchase Order found for Order Code {order.order_number}"
    return order.purchase_order, None
