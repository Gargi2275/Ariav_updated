"""Parser and reconciliation for the one-time party statement workbook."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

MONEY_ZERO = Decimal("0.00")
DATE_RE = re.compile(r"(\d{2}\.\d{2}\.\d{4})")
CODE_RE = re.compile(r"Customer\s*Code\s*:\s*([^\s]+)", re.IGNORECASE)
LEDGER_CODE_RE = re.compile(r"Ledger\s*Code\s*:\s*([^\s]+)", re.IGNORECASE)
PERIOD_RE = re.compile(r"FROM\s+(\d{2}\.\d{2}\.\d{4})\s+TO\s+(\d{2}\.\d{2}\.\d{4})", re.IGNORECASE)


def cell_text(value) -> str:
    return "" if value is None else str(value).strip()


def parse_date(value: str, row_number: int, label: str, errors: list[str]) -> date | None:
    raw = cell_text(value)
    if not raw:
        return None
    try:
        return datetime.strptime(raw, "%d.%m.%Y").date()
    except ValueError:
        errors.append(f"Excel row {row_number}: invalid {label} date {raw!r}.")
    return None


def parse_amount(value, row_number: int, label: str, errors: list[str]) -> Decimal:
    raw = cell_text(value).replace(",", "")
    if not raw:
        return MONEY_ZERO
    trailing_minus = raw.endswith("-")
    if trailing_minus:
        raw = raw[:-1].strip()
    try:
        amount = Decimal(raw)
    except InvalidOperation:
        errors.append(f"Excel row {row_number}: invalid {label} amount {value!r}.")
        return MONEY_ZERO
    return amount.copy_abs().quantize(MONEY_ZERO) if trailing_minus else amount.quantize(MONEY_ZERO)


def signed_amount(amount: Decimal, dc: str) -> Decimal:
    return amount if dc == "DR" else -amount


@dataclass
class EntryData:
    excel_row: int
    line_no: int
    post_date: date
    doc_type: str
    doc_no: str
    text: str
    debit: Decimal
    credit: Decimal
    balance_amount: Decimal
    balance_dc: str
    reference: str
    tds_amount: Decimal
    due_date: date | None
    agent_code: str


@dataclass
class StatementData:
    customer_code: str
    customer_name: str
    period_from: date | None
    period_to: date | None
    ledger_code: str
    ledger_description: str
    opening_date: date | None
    opening_amount: Decimal
    opening_dc: str
    closing_date: date | None
    total_debit: Decimal
    total_credit: Decimal
    closing_amount: Decimal
    closing_dc: str
    entries: list[EntryData] = field(default_factory=list)
    start_row: int = 0


@dataclass
class ParsedWorkbook:
    statements: list[StatementData]
    errors: list[str]


def _label_value(text: str, label: str) -> str:
    match = re.search(rf"{re.escape(label)}\s*:\s*(.*?)(?=\s{{2,}}[A-Za-z ]+\s*:|$)", text, re.IGNORECASE)
    return match.group(1).strip() if match else ""


def _parse_block(rows: list[tuple[int, list]], start_index: int, period_from: date | None, period_to: date | None, errors: list[str]) -> tuple[StatementData, int]:
    row_number, values = rows[start_index]
    first = cell_text(values[0])
    code_match = CODE_RE.search(first)
    ledger_match = LEDGER_CODE_RE.search(first)
    code = code_match.group(1).strip() if code_match else ""
    ledger_code = ledger_match.group(1).strip() if ledger_match else ""
    name = _label_value(first, "Customer Name") or _label_value(cell_text(rows[start_index + 1][1][0]), "Customer Name")
    details = [cell_text(rows[i][1][0]) for i in range(start_index, min(start_index + 6, len(rows)))]
    name = next((_label_value(text, "Customer Name") for text in details if _label_value(text, "Customer Name")), name)
    description = next((_label_value(text, "Ledger Description") for text in details if _label_value(text, "Ledger Description")), "")
    statement = StatementData(
        customer_code=code,
        customer_name=name,
        period_from=period_from,
        period_to=period_to,
        ledger_code=ledger_code,
        ledger_description=description,
        opening_date=None,
        opening_amount=MONEY_ZERO,
        opening_dc="DR",
        closing_date=None,
        total_debit=MONEY_ZERO,
        total_credit=MONEY_ZERO,
        closing_amount=MONEY_ZERO,
        closing_dc="DR",
        start_row=row_number,
    )
    header_index = None
    opening_seen = False
    closing_seen = False
    line_no = 0
    index = start_index + 1
    while index < len(rows):
        current_row, current = rows[index]
        cells = [cell_text(value) for value in current]
        first_cell = cells[0] if cells else ""
        row_text = " ".join(cells)
        if index > start_index and CODE_RE.search(first_cell):
            errors.append(f"Customer {code or '<missing>'}, Excel row {row_number}: missing closing balance row.")
            return statement, index
        if "POST DATE" in first_cell.upper() and "DOC. TYPE" in row_text.upper():
            header_index = index
            index += 1
            continue
        if "OPENING BALANCE" in row_text.upper():
            statement.opening_date = parse_date(DATE_RE.search(row_text).group(1) if DATE_RE.search(row_text) else "", current_row, "opening", errors)
            statement.opening_amount = parse_amount(current[6] if len(current) > 6 else "", current_row, "opening balance", errors)
            statement.opening_dc = (cells[7] if len(cells) > 7 else "").upper()
            opening_seen = True
            index += 1
            continue
        if "CLOSING BALANCE" in row_text.upper():
            statement.closing_date = parse_date(DATE_RE.search(row_text).group(1) if DATE_RE.search(row_text) else "", current_row, "closing", errors)
            statement.total_debit = parse_amount(current[4] if len(current) > 4 else "", current_row, "closing debit total", errors)
            statement.total_credit = parse_amount(current[5] if len(current) > 5 else "", current_row, "closing credit total", errors)
            statement.closing_amount = parse_amount(current[6] if len(current) > 6 else "", current_row, "closing balance", errors)
            statement.closing_dc = (cells[7] if len(cells) > 7 else "").upper()
            closing_seen = True
            return statement, index + 1
        if header_index is not None and first_cell:
            post_date = parse_date(first_cell, current_row, "posting", errors)
            if post_date is None:
                errors.append(f"Customer {code or '<missing>'}, Excel row {current_row}: invalid transaction row.")
                index += 1
                continue
            doc_type = cells[1] if len(cells) > 1 else ""
            doc_no = cells[2] if len(cells) > 2 else ""
            if not doc_type or not doc_no:
                errors.append(f"Customer {code or '<missing>'}, Excel row {current_row}: transaction type and number are required.")
            line_no += 1
            statement.entries.append(
                EntryData(
                    excel_row=current_row,
                    line_no=line_no,
                    post_date=post_date,
                    doc_type=doc_type,
                    doc_no=doc_no,
                    text=cells[3] if len(cells) > 3 else "",
                    debit=parse_amount(current[4] if len(current) > 4 else "", current_row, "debit", errors),
                    credit=parse_amount(current[5] if len(current) > 5 else "", current_row, "credit", errors),
                    balance_amount=parse_amount(current[6] if len(current) > 6 else "", current_row, "balance", errors),
                    balance_dc=(cells[7] if len(cells) > 7 else "").upper(),
                    reference=cells[8] if len(cells) > 8 else "",
                    tds_amount=parse_amount(current[9] if len(current) > 9 else "", current_row, "TDS", errors),
                    due_date=parse_date(cells[10] if len(cells) > 10 else "", current_row, "due", errors),
                    agent_code=cells[11] if len(cells) > 11 else "",
                )
            )
        index += 1
    if not opening_seen:
        errors.append(f"Customer {code or '<missing>'}, Excel row {row_number}: opening balance row not found.")
    if not closing_seen:
        errors.append(f"Customer {code or '<missing>'}, Excel row {row_number}: closing balance row not found.")
    if header_index is None:
        errors.append(f"Customer {code or '<missing>'}, Excel row {row_number}: transaction header row not found.")
    return statement, index


def parse_workbook(path: str | Path) -> ParsedWorkbook:
    workbook = load_workbook(Path(path), data_only=True, read_only=True)
    errors: list[str] = []
    if len(workbook.sheetnames) != 1:
        errors.append(f"Workbook must contain exactly one sheet; found {workbook.sheetnames}.")
    sheet = workbook[workbook.sheetnames[0]]
    rows = [(index, list(values)) for index, values in enumerate(sheet.iter_rows(values_only=True), start=1)]
    starts = [index for index, values in enumerate(rows) if CODE_RE.search(cell_text(values[1][0]))]
    statements: list[StatementData] = []
    for position, start in enumerate(starts):
        preceding = rows[max(0, start - 4):start]
        period_from = period_to = None
        for _, values in reversed(preceding):
            match = PERIOD_RE.search(cell_text(values[0]))
            if match:
                period_from = parse_date(match.group(1), rows[start][0], "period from", errors)
                period_to = parse_date(match.group(2), rows[start][0], "period to", errors)
                break
        statement, _ = _parse_block(rows, start, period_from, period_to, errors)
        statements.append(statement)
    workbook.close()
    if not statements:
        errors.append("No customer blocks found.")
    return ParsedWorkbook(statements=statements, errors=errors)


def validate_statements(statements: list[StatementData]) -> list[str]:
    errors: list[str] = []
    periods = {(statement.period_from, statement.period_to) for statement in statements}
    if len(periods) > 1:
        errors.append("Statements contain inconsistent reporting periods.")
    seen_customers: set[str] = set()
    for statement in statements:
        code = statement.customer_code or "<missing>"
        if not statement.customer_code:
            errors.append(f"Customer {code}, Excel row {statement.start_row}: customer code is missing.")
        elif statement.customer_code.upper() in seen_customers:
            errors.append(f"Customer {code}, Excel row {statement.start_row}: duplicate customer block.")
        seen_customers.add(statement.customer_code.upper())
        if statement.period_from is None or statement.period_to is None:
            errors.append(f"Customer {code}, Excel row {statement.start_row}: statement period is missing.")
        if statement.opening_dc not in {"DR", "CR"}:
            errors.append(f"Customer {code}, Excel row {statement.start_row}: opening DR/CR is invalid.")
        if statement.closing_dc not in {"DR", "CR"}:
            errors.append(f"Customer {code}, Excel row {statement.start_row}: closing DR/CR is invalid.")
        seen: set[tuple[str, str]] = set()
        running = signed_amount(statement.opening_amount, statement.opening_dc)
        debit_total = MONEY_ZERO
        credit_total = MONEY_ZERO
        for entry in statement.entries:
            key = (entry.doc_type, entry.doc_no)
            if key in seen:
                errors.append(f"Customer {code}, Excel row {entry.excel_row}: duplicate document {entry.doc_type}/{entry.doc_no}.")
            seen.add(key)
            debit_total += entry.debit
            credit_total += entry.credit
            running = (running + entry.debit - entry.credit).quantize(MONEY_ZERO)
            file_running = signed_amount(entry.balance_amount, entry.balance_dc)
            if running != file_running:
                errors.append(
                    f"Customer {code}, Excel row {entry.excel_row}: running balance {file_running} "
                    f"does not match calculated {running}."
                )
        if debit_total != statement.total_debit:
            errors.append(f"Customer {code}, Excel row {statement.start_row}: debit total {debit_total} != closing total {statement.total_debit}.")
        if credit_total != statement.total_credit:
            errors.append(f"Customer {code}, Excel row {statement.start_row}: credit total {credit_total} != closing total {statement.total_credit}.")
        expected_closing = (signed_amount(statement.opening_amount, statement.opening_dc) + debit_total - credit_total).quantize(MONEY_ZERO)
        actual_closing = signed_amount(statement.closing_amount, statement.closing_dc)
        if expected_closing != actual_closing:
            errors.append(f"Customer {code}, Excel row {statement.start_row}: closing balance {actual_closing} != calculated {expected_closing}.")
    return errors
