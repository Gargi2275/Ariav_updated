"""One-off bulk import of Aditi Entity customers from cleaned Excel."""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from customers.models import Customer, CustomerEntity
from entities.models import Entity

PLACEHOLDER_PIN = "000000"
DUPLICATE_PAIRS_HINT = (
    ("TARAK TEXTORIUM", ("700777", "703122")),
    ("BHAGWAN SUITING SHIRTING", ("700799", "703236")),
    ("CONNOISSEUR", ("702407", "702411")),
)


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _readable_place(value: str) -> str:
    value = (value or "").strip()
    if value and value.isupper():
        return value.title()
    return value


def _parse_city(raw: str) -> tuple[str, str, bool]:
    raw = (raw or "").strip()
    if " - " in raw:
        city_name, pincode_raw = raw.split(" - ", 1)
    else:
        city_name, pincode_raw = raw, ""
    city = _readable_place(city_name)
    pin = pincode_raw.strip()
    blanked = pin == PLACEHOLDER_PIN
    if blanked or not pin:
        pin = ""
    return city, pin, blanked


def _code_str(value) -> str:
    raw = _cell_str(value)
    return raw


class Command(BaseCommand):
    help = (
        "Import real Aditi Entity customers from aditi_textile_customers_clean.xlsx. "
        "Use --dry-run first; the live import is one transaction and rolls back on any failure."
    )

    def add_arguments(self, parser):
        parser.add_argument("xlsx_path", help="Path to aditi_textile_customers_clean.xlsx")
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Validate and report only; do not write to the database.",
        )

    def handle(self, *args, **options):
        path = Path(options["xlsx_path"]).expanduser().resolve()
        if not path.exists():
            raise CommandError(f"File not found: {path}")

        aditi = (
            Entity.objects.filter(entity_name__iexact="Aditi").first()
            or Entity.objects.filter(short_code__iexact="ADT").first()
        )
        if aditi is None:
            raise CommandError("Aditi Entity not found (name 'Aditi' or short_code 'ADT').")

        rows, parse_notes = self._read_rows(path)
        report = self._validate(rows, aditi)

        self._print_report(path, aditi, rows, report, parse_notes, dry_run=options["dry_run"])

        if report["collisions"] or report["invalid"]:
            raise CommandError(
                "Import blocked: collisions or validation failures. See report above."
            )

        if options["dry_run"]:
            self.stdout.write(self.style.WARNING("DRY RUN — no records written."))
            return

        created_customers, created_links = self._import(rows, aditi)
        self.stdout.write(
            self.style.SUCCESS(
                f"Imported {created_customers} customers and {created_links} Aditi primary entity links."
            )
        )
        self.stdout.write(
            f"Placeholder pincodes blanked: {report['blanked_pincodes']}"
        )
        self._print_duplicate_pairs(report["exact_duplicate_groups"])

    def _read_rows(self, path: Path) -> tuple[list[dict], list[str]]:
        wb = load_workbook(path, data_only=True, read_only=True)
        if "Customers" not in wb.sheetnames:
            raise CommandError(f"Sheet 'Customers' not found. Sheets: {wb.sheetnames}")
        ws = wb["Customers"]
        notes: list[str] = []
        rows: list[dict] = []
        header = None
        for idx, raw in enumerate(ws.iter_rows(values_only=True), start=1):
            values = list(raw[:5])
            if idx == 1:
                header = [_cell_str(v).lower() for v in values]
                expected = ["code", "name", "address", "city", "state"]
                if header != expected:
                    raise CommandError(f"Unexpected header {header}; expected {expected}.")
                continue
            if all(_cell_str(v) == "" for v in values):
                continue
            code = _code_str(values[0])
            name = _cell_str(values[1])
            address = _cell_str(values[2])
            city, pincode, blanked = _parse_city(_cell_str(values[3]))
            state = _cell_str(values[4])
            rows.append(
                {
                    "excel_row": idx,
                    "customer_code": code,
                    "customer_name": name,
                    "address_line_1": address,
                    "city": city,
                    "pincode": pincode,
                    "pincode_blanked": blanked,
                    "state": state,
                    "country": "India",
                    "customer_type": Customer.CustomerType.COMPANY,
                    "status": Customer.Status.ACTIVE,
                }
            )
        wb.close()
        if len(rows) != 263:
            notes.append(f"Expected 263 data rows; read {len(rows)}.")
        return rows, notes

    def _validate(self, rows: list[dict], aditi: Entity) -> dict:
        existing_ci = {
            c.upper(): c
            for c in Customer.objects.values_list("customer_code", flat=True)
        }

        collisions = []
        invalid = []
        seen_in_file: dict[str, int] = {}
        for row in rows:
            code = row["customer_code"]
            name = row["customer_name"]
            errors = []
            if not code:
                errors.append("customer_code is required")
            elif len(code) > 30:
                errors.append("customer_code exceeds 30 characters")
            if not name:
                errors.append("customer_name is required")
            elif len(name) > 200:
                errors.append("customer_name exceeds 200 characters")
            if len(row["address_line_1"]) > 200:
                errors.append("address_line_1 exceeds 200 characters")
            if len(row["city"]) > 80:
                errors.append("city exceeds 80 characters")
            if len(row["state"]) > 80:
                errors.append("state exceeds 80 characters")
            if len(row["pincode"]) > 12:
                errors.append("pincode exceeds 12 characters")
            if code in seen_in_file:
                errors.append(f"duplicate code in file (also row {seen_in_file[code]})")
            else:
                seen_in_file[code] = row["excel_row"]
            hit = existing_ci.get(code.upper())
            if hit:
                collisions.append(
                    {
                        "excel_row": row["excel_row"],
                        "customer_code": code,
                        "existing_code": hit,
                        "customer_name": name,
                    }
                )
            if errors:
                invalid.append({"excel_row": row["excel_row"], "customer_code": code, "errors": errors})

        groups: dict[tuple[str, str], list[dict]] = defaultdict(list)
        for row in rows:
            key = (row["customer_name"].casefold(), row["address_line_1"].casefold())
            groups[key].append(row)
        exact_duplicate_groups = [
            sorted(members, key=lambda r: r["customer_code"])
            for members in groups.values()
            if len(members) > 1
        ]
        exact_duplicate_groups.sort(key=lambda g: g[0]["customer_name"])

        return {
            "collisions": collisions,
            "invalid": invalid,
            "blanked_pincodes": sum(1 for r in rows if r["pincode_blanked"]),
            "exact_duplicate_groups": exact_duplicate_groups,
            "existing_customer_count": Customer.objects.count(),
            "existing_aditi_links": CustomerEntity.objects.filter(
                entity=aditi, primary_entity=True
            ).count(),
        }

    def _print_report(self, path, aditi, rows, report, notes, dry_run: bool):
        mode = "DRY RUN" if dry_run else "IMPORT"
        self.stdout.write(self.style.MIGRATE_HEADING(f"=== {mode}: Aditi customer import ==="))
        self.stdout.write(f"File: {path}")
        self.stdout.write(f"Aditi entity: id={aditi.id} {aditi.short_code} · {aditi.entity_name}")
        self.stdout.write(f"Existing customers in DB: {report['existing_customer_count']}")
        self.stdout.write(f"Rows to import: {len(rows)}")
        self.stdout.write(f"CustomerEntity links that would be created: {len(rows)} (Aditi, primary)")
        self.stdout.write(f"Placeholder pincodes 000000 blanked: {report['blanked_pincodes']}")
        for note in notes:
            self.stdout.write(self.style.WARNING(f"Note: {note}"))

        if report["collisions"]:
            self.stdout.write(self.style.ERROR(f"Code collisions with existing customers: {len(report['collisions'])}"))
            for item in report["collisions"]:
                self.stdout.write(
                    f"  row {item['excel_row']}: file {item['customer_code']} "
                    f"collides with existing {item['existing_code']} ({item['customer_name']})"
                )
        else:
            self.stdout.write(self.style.SUCCESS("Code collisions with existing customers: 0"))

        if report["invalid"]:
            self.stdout.write(self.style.ERROR(f"Rows failing required-field validation: {len(report['invalid'])}"))
            for item in report["invalid"]:
                self.stdout.write(
                    f"  row {item['excel_row']} code={item['customer_code']}: {'; '.join(item['errors'])}"
                )
        else:
            self.stdout.write(self.style.SUCCESS("Required-field validation failures: 0"))

        self._print_duplicate_pairs(report["exact_duplicate_groups"])

    def _print_duplicate_pairs(self, groups: list[list[dict]]):
        self.stdout.write("")
        self.stdout.write(
            self.style.WARNING(
                f"Exact duplicate pairs (same name AND same address, different codes) — "
                f"{len(groups)} pair(s). Both records will be imported as separate ledger accounts."
            )
        )
        if not groups:
            self.stdout.write("  (none found)")
            return
        for group in groups:
            name = group[0]["customer_name"]
            addr = group[0]["address_line_1"]
            codes = "/".join(r["customer_code"] for r in group)
            self.stdout.write(f"  FLAG {name} ({codes})")
            self.stdout.write(f"       address: {addr}")
        hinted = {name.upper() for name, _ in DUPLICATE_PAIRS_HINT}
        found = {g[0]["customer_name"].upper() for g in groups}
        missing_hint = hinted - found
        extra = found - hinted
        if missing_hint:
            self.stdout.write(self.style.WARNING(f"  Expected pairs not found: {sorted(missing_hint)}"))
        if extra:
            self.stdout.write(self.style.WARNING(f"  Extra exact-duplicate groups beyond the known 3: {sorted(extra)}"))

    @transaction.atomic
    def _import(self, rows: list[dict], aditi: Entity) -> tuple[int, int]:
        customers = []
        for row in rows:
            customers.append(
                Customer(
                    customer_code=row["customer_code"],
                    customer_name=row["customer_name"],
                    customer_type=row["customer_type"],
                    status=row["status"],
                    address_line_1=row["address_line_1"],
                    city=row["city"],
                    pincode=row["pincode"],
                    state=row["state"],
                    country=row["country"],
                )
            )
        created = Customer.objects.bulk_create(customers, batch_size=100)
        # Re-fetch with PKs (MySQL bulk_create may not set them on all versions).
        codes = [c.customer_code for c in created]
        saved = {c.customer_code: c for c in Customer.objects.filter(customer_code__in=codes)}
        if len(saved) != len(rows):
            raise CommandError(
                f"Import aborted: expected {len(rows)} customers after insert, found {len(saved)}."
            )
        links = [
            CustomerEntity(customer=saved[row["customer_code"]], entity=aditi, primary_entity=True)
            for row in rows
        ]
        CustomerEntity.objects.bulk_create(links, batch_size=100)
        return len(saved), len(links)
